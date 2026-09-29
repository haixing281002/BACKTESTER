"""compute_decile_membership / run_decile_backtests / summarize_deciles /
write_decile_membership_log -- the full NIFTY 500 decile deep-dive, not
just the top/bottom decile pair."""
import numpy as np
import pandas as pd
import pytest

from universal_backtester.decile_analysis import (
    compute_decile_membership, run_decile_backtests, summarize_deciles,
    write_decile_membership_log, write_decile_summary_workbook,
)


def _fake_universe(n_assets=100, n_days=500, seed=0):
    rng = np.random.default_rng(seed)
    dates = pd.bdate_range("2020-01-01", periods=n_days)
    assets = list(range(1, n_assets + 1))
    prices = pd.DataFrame(
        100 * (1 + rng.normal(0.0004, 0.012, (n_days, n_assets))).cumprod(axis=0),
        index=dates, columns=assets)
    mcap = pd.DataFrame(np.tile(np.arange(1, n_assets + 1, dtype=float), (n_days, 1)),
                        index=dates, columns=assets)
    eligible = pd.DataFrame(True, index=dates, columns=assets)
    return prices, mcap, eligible, assets


def test_decile_membership_splits_into_equal_count_buckets_by_rank():
    _, mcap, eligible, assets = _fake_universe(n_assets=100, n_days=5)
    masks = compute_decile_membership(mcap, eligible, n_deciles=10)
    assert set(masks) == set(range(1, 11))
    d = mcap.index[0]
    for k in range(1, 11):
        assert masks[k].loc[d].sum() == 10
    # decile 1 must hold the 10 smallest-mcap codes (1..10), decile 10 the largest (91..100)
    assert set(mcap.columns[masks[1].loc[d]]) == set(range(1, 11))
    assert set(mcap.columns[masks[10].loc[d]]) == set(range(91, 101))


def test_decile_membership_excludes_ineligible_names_from_every_bucket():
    _, mcap, eligible, assets = _fake_universe(n_assets=20, n_days=3)
    eligible = eligible.copy()
    eligible.iloc[:, :5] = False   # first 5 (smallest mcap) names are not eligible
    masks = compute_decile_membership(mcap, eligible, n_deciles=10)
    d = mcap.index[0]
    for k in range(1, 11):
        chosen = mcap.columns[masks[k].loc[d]]
        assert not any(c in mcap.columns[:5] for c in chosen)


def test_run_decile_backtests_produces_one_result_per_decile():
    prices, mcap, eligible, assets = _fake_universe(n_assets=50, n_days=400)
    masks = compute_decile_membership(mcap, eligible, n_deciles=10)
    results = run_decile_backtests(prices, assets, mcap, masks, spread_bps=30,
                                   lag_days=1, rebalance="monthly", warmup=30,
                                   min_names=2)
    assert set(results) == set(range(1, 11))
    for k, r in results.items():
        assert len(r.value.dropna()) > 10

    summary = summarize_deciles(results)
    assert list(summary["decile"]) == list(range(1, 11))
    assert {"cagr", "vol", "sharpe", "max_dd"}.issubset(summary.columns)


def test_decile_membership_log_names_every_stock_and_its_decile(tmp_path):
    _, mcap, eligible, assets = _fake_universe(n_assets=30, n_days=10)
    masks = compute_decile_membership(mcap, eligible, n_deciles=10)
    dates = mcap.index[:3]
    name_lookup = {a: f"Company {a}" for a in assets}
    log = write_decile_membership_log(str(tmp_path / "decile_membership.csv"),
                                      masks, dates, mcap, name_lookup=name_lookup)
    assert set(log["date"].unique()) == set(dates)
    # every eligible name on a date appears exactly once across all deciles that date
    for d in dates:
        day = log[log["date"] == d]
        assert len(day) == 30
        assert day["decile"].value_counts().eq(3).all()
    assert (log["name"] == log["accord_code"].map(name_lookup)).all()


def test_decile_workbook_embeds_every_chart_as_its_own_sheet(tmp_path):
    import openpyxl
    summary = pd.DataFrame({"decile": [1, 2], "start": ["2020-01-01"] * 2, "end": ["2020-12-31"] * 2,
                            "n_obs": [250, 250], "cagr": [0.1, 0.2], "vol": [0.15, 0.18],
                            "sharpe": [0.5, 0.8], "max_dd": [-0.2, -0.15]})
    # A minimal real PNG so openpyxl's Image loader has something to actually decode.
    chart1 = tmp_path / "chart1.png"
    chart2 = tmp_path / "chart2.png"
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    for p in (chart1, chart2):
        fig, ax = plt.subplots()
        ax.plot([1, 2, 3], [1, 2, 3])
        fig.savefig(p)
        plt.close(fig)

    out = tmp_path / "decile_summary.xlsx"
    write_decile_summary_workbook(str(out), summary, chart_paths=[str(chart1), str(chart2)])

    wb = openpyxl.load_workbook(str(out))
    assert wb.sheetnames[0] == "Decile Summary"
    assert len(wb.sheetnames) == 3
    for sheet in wb.sheetnames[1:]:
        assert len(wb[sheet]._images) == 1


def test_decile_workbook_notes_a_missing_chart_file_rather_than_silently_dropping_it(tmp_path):
    import openpyxl
    summary = pd.DataFrame({"decile": [1], "start": ["2020-01-01"], "end": ["2020-12-31"],
                            "n_obs": [250], "cagr": [0.1], "vol": [0.15], "sharpe": [0.5],
                            "max_dd": [-0.2]})
    out = tmp_path / "decile_summary.xlsx"
    write_decile_summary_workbook(str(out), summary, chart_paths=[str(tmp_path / "does_not_exist.png")])

    wb = openpyxl.load_workbook(str(out))
    missing_sheet = wb[wb.sheetnames[1]]
    assert "not found" in missing_sheet["A1"].value.lower()
