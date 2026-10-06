"""Tests for scripts/reconcile_accord_vs_nse.py, built against synthetic
Accord files (same convention as test_universal_backtester_accord_data.py)
and synthetic bhavcopy pickles (same convention as
test_nse_bhavcopy_ingest.py) -- this sandbox has no real NSE cache and the
real Accord files, where present, are local-only.

The scenario: two symbols, RELIANCE and WIPRO, with matching Accord Codes.
RELIANCE's two sources agree closely (should report "ok"). WIPRO's two
sources are deliberately built to disagree by a large, sustained amount
(should report "REVIEW") -- this is the actual failure mode the script
exists to catch: a corporate action one source adjusted for and the other
didn't, or a vendor restatement nobody told the pipeline about.
"""
import numpy as np
import openpyxl
import pandas as pd

from scripts.reconcile_accord_vs_nse import bhavcopy_symbol_map, reconcile, run
from ros.data.nse_bhavcopy_ingest import build_adjusted_prices, combine_cache, dedupe_rows, link_isins
from universal_backtester.accord_data import (
    build_accord_ticker_bridge, load_accord_monthly_universe, load_accord_price_panel,
)


def _make_accord_price_file(path, codes, dates, closes):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["NDP_Date", *codes])
    for d, row in zip(dates, closes):
        ws.append([d.strftime("%Y-%m-%d"), *row])
    wb.save(path)


def _make_accord_universe_file(path, code_to_symbol):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "31-Jan-2023"
    ws.append(["Accord Code", "Company Name", "NDP_Date", "NDP_Close", "NDP_Mcap", "NSE_symbol", "Market Rank"])
    for i, (code, sym) in enumerate(code_to_symbol.items(), start=1):
        ws.append([code, sym, "2023-01-31", 100.0, 1000.0, sym, i])
    wb.save(path)


def _row(date, symbol, isin, close, prev, val=50_00_00_000):
    return dict(DATE=pd.Timestamp(date), SYMBOL=symbol, SERIES="EQ", ISIN=isin,
                OPEN=close, HIGH=close, LOW=close, CLOSE=close, PREV=prev,
                QTY=10_000, VAL=val, TRADES=500)


def _make_bhavcopy(days, rel_prices, wip_prices=None):
    """RELIANCE: pass the SAME price path Accord uses (near-identical two
    sources, as a matching pair should be). WIPRO: pass a price path built
    independently of Accord's, so it diverges (simulating an unreconciled
    disagreement) -- or omit it to use a flat default."""
    rows = []
    if wip_prices is None:
        wip_prices = np.full(len(days), 400.0)
    rel_prev = rel_prices[0]
    wip_prev = wip_prices[0]
    for d, rel_price, wip_price in zip(days, rel_prices, wip_prices):
        rows.append(_row(d, "RELIANCE", "INERELI01018", rel_price, rel_prev))
        rel_prev = rel_price
        rows.append(_row(d, "WIPRO", "INEWIP001011", wip_price, wip_prev))
        wip_prev = wip_price
    return pd.DataFrame(rows)


def _write_cache(df, cache_dir):
    for d, g in df.groupby("DATE"):
        g.to_pickle(cache_dir / f"{d:%Y%m%d}.pkl")


def test_matching_series_is_ok_and_diverging_series_is_flagged(tmp_path):
    days = pd.bdate_range("2023-01-02", periods=120)
    rng = np.random.default_rng(0)

    # Accord side: RELIANCE tracks the same gentle drift as bhavcopy will;
    # WIPRO drifts flat (so it diverges hard from bhavcopy's 1%/day drift).
    rel_accord = 2500.0 * np.cumprod(1.0005 + rng.normal(0, 0.0005, len(days)))
    wip_accord = 400.0 * np.cumprod(1.0 + rng.normal(0, 0.0005, len(days)))
    closes = np.column_stack([rel_accord, wip_accord])

    price_path = tmp_path / "accord_price.xlsx"
    _make_accord_price_file(price_path, [111, 222], days, closes)

    universe_path = tmp_path / "accord_universe.xlsx"
    _make_accord_universe_file(universe_path, {111: "RELIANCE", 222: "WIPRO"})

    cache_dir = tmp_path / "cache"
    cache_dir.mkdir()
    # RELIANCE: bhavcopy uses Accord's own series (near-identical sources).
    # WIPRO: bhavcopy drifts at 1%/day, independent of Accord's flat series.
    wip_bhav = 400.0 * np.cumprod(np.full(len(days), 1.01))
    bhav = _make_bhavcopy(days, rel_accord, wip_bhav)
    _write_cache(bhav, cache_dir)

    out_path = tmp_path / "report.csv"
    report = run(str(price_path), str(universe_path), str(cache_dir), str(out_path),
                min_common_days=40)

    assert out_path.exists()
    assert set(report["symbol"]) == {"RELIANCE", "WIPRO"}

    rel_row = report[report["symbol"] == "RELIANCE"].iloc[0]
    assert rel_row["status"] == "ok"
    assert rel_row["return_corr"] > 0.9

    wip_row = report[report["symbol"] == "WIPRO"].iloc[0]
    assert wip_row["status"] == "REVIEW"
    assert wip_row["max_abs_rebased_dev_pct"] > 10.0


def test_symbol_ever_changed_is_noted_not_hidden(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "31-Jan-2023"
    ws.append(["Accord Code", "Company Name", "NDP_Date", "NDP_Close", "NDP_Mcap", "NSE_symbol", "Market Rank"])
    ws.append([111, "OldName", "2023-01-31", 100.0, 1000.0, "OLDSYM", 1])
    ws2 = wb.create_sheet("28-Feb-2023")
    ws2.append(["Accord Code", "Company Name", "NDP_Date", "NDP_Close", "NDP_Mcap", "NSE_symbol", "Market Rank"])
    ws2.append([111, "NewName", "2023-02-28", 101.0, 1010.0, "NEWSYM", 1])
    path = tmp_path / "universe.xlsx"
    wb.save(path)

    universe_df, _ = load_accord_monthly_universe(str(path))
    bridge, _ = build_accord_ticker_bridge(universe_df)
    row = bridge[bridge["accord_code"] == 111].iloc[0]
    assert row["symbol_ever_changed"]
    assert row["nse_symbol"] == "NEWSYM"
    assert "OLDSYM" in row["prior_symbols"]


def test_symbol_to_code_map_is_correct_when_some_codes_have_no_symbol(tmp_path):
    """Regression: the bridge carries many rows with accord_code but a NULL
    nse_symbol (a code never matched to a symbol in the source file). Building
    the symbol->code map by calling .dropna() on the two columns SEPARATELY
    before zipping silently pairs each symbol with the wrong code once the
    two filtered columns have different lengths -- caught by running this
    script against the real Accord files, where it produced a RELIANCE row
    actually pointing at a different company's accord_code."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "31-Jan-2023"
    ws.append(["Accord Code", "Company Name", "NDP_Date", "NDP_Close", "NDP_Mcap", "NSE_symbol", "Market Rank"])
    # code 50 has NO symbol (the common real-world case); code 111 (RELIANCE)
    # comes right after it. A positional zip after separate dropna() shifts
    # every subsequent row by one and pairs RELIANCE with the wrong code.
    ws.append([50, "Unknown Co.", "2023-01-31", 10.0, 100.0, None, 99])
    ws.append([111, "Reliance", "2023-01-31", 2500.0, 900000.0, "RELIANCE", 1])
    path = tmp_path / "universe.xlsx"
    wb.save(path)

    universe_df, _ = load_accord_monthly_universe(str(path))
    bridge, _ = build_accord_ticker_bridge(universe_df)

    days = pd.bdate_range("2023-01-02", periods=60)
    closes = np.column_stack([np.linspace(2500.0, 2600.0, len(days))])
    price_path = tmp_path / "accord_price.xlsx"
    _make_accord_price_file(price_path, [111], days, closes)

    cache_dir = tmp_path / "cache"
    cache_dir.mkdir()
    bhav = _make_bhavcopy(days, np.linspace(2500.0, 2600.0, len(days)))
    _write_cache(bhav, cache_dir)

    accord_price, _ = load_accord_price_panel(str(price_path))
    nse_raw = dedupe_rows(combine_cache(str(cache_dir)))
    links = link_isins(nse_raw)
    adjusted = build_adjusted_prices(nse_raw, links)
    nse_sym = bhavcopy_symbol_map(nse_raw, links)

    report = reconcile(accord_price, bridge, adjusted.panel, nse_sym, min_common_days=40)
    row = report[report["symbol"] == "RELIANCE"].iloc[0]
    assert row["accord_code"] == 111
    assert row["status"] == "ok"


def test_too_few_common_days_is_reported_not_silently_dropped(tmp_path):
    days = pd.bdate_range("2023-01-02", periods=10)  # below default min_common_days=40
    closes = np.column_stack([np.linspace(100, 110, len(days))])

    price_path = tmp_path / "accord_price.xlsx"
    _make_accord_price_file(price_path, [111], days, closes)

    universe_path = tmp_path / "accord_universe.xlsx"
    _make_accord_universe_file(universe_path, {111: "RELIANCE"})

    cache_dir = tmp_path / "cache"
    cache_dir.mkdir()
    bhav = _make_bhavcopy(days, np.full(len(days), 2500.0))
    _write_cache(bhav, cache_dir)

    out_path = tmp_path / "report.csv"
    report = run(str(price_path), str(universe_path), str(cache_dir), str(out_path))

    rel_row = report[report["symbol"] == "RELIANCE"].iloc[0]
    assert rel_row["status"] == "too_few_common_days"
