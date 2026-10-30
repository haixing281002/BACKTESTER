"""NIFTY 500 membership, point in time, from the holdings workbook (ros/data/sources.py: NIFTY500_HOLDINGS_XLSX).

THE FILE
    "20261022 Nifty 500 Historical Holdings since 2013.xlsx". Its "Bifurcation" sheet lists, for every quarter-end
    from 2012-03-31 to the latest pull, the Bloomberg members of the NSE 500 index with the ISIN and NSE symbol
    looked up from the workbook's own security master, and a flag for which sub-index each name sat in:
    Nifty 50, Nifty Next 50, Nifty 100, Nifty Midcap 150, Nifty Smallcap 250. The workbook is read with the
    cached values Excel saved, so it must have been calculated in Excel before it is dropped here.

HOW A SNAPSHOT BECOMES MEMBERSHIP
    A snapshot dated d says who was in the index at the close of d. A name is treated as a member from d
    (inclusive) until the first later snapshot that no longer lists it (exclusive). Nothing between snapshots is
    interpolated: a name that joined at an NSE semi-annual review in March is a member from the 31 March
    snapshot, which is the first date the list was observable. Snapshots are quarterly, NSE reviews are
    semi-annual (effective end-March and end-September), so the intermediate quarter-ends normally repeat the
    previous list.

WHAT IS CHECKED, AND WHAT IS LEFT OUT (all reported by `snapshot_summary()` and `provenance()`)
    * A snapshot with fewer than MIN_COMPLETE names is an incomplete pull (the file has three: 2017-09-30,
      2018-06-30, 2018-09-30). It is SKIPPED by default and the previous complete snapshot carries forward,
      because a short list would silently drop ~100-200 names from the universe for a quarter.
    * Members whose Bloomberg id never resolved to an ISIN (delisted or merged before the master was exported,
      e.g. "1909315D IS Equity") are unmapped: they cannot be priced and are left out. 40 names in 2012,
      falling to 0 by 2026. This IS survivorship bias at the margin and every result built on this file must
      say so; the counts are in `snapshot_summary()`.
    * After the last snapshot the list is carried forward. `staleness(end_date)` says whether an NSE review
      (end-March / end-September) has happened since the last snapshot; a run that extends past that date must
      report it. To bring the file up to date without touching the workbook, drop a CSV named YYYY-MM-DD.csv with
      an `isin` column (optionally `symbol`, `company_name`) into data/raw/nifty500_updates/; it is read as one
      more snapshot.

    from ros.data.nifty500 import members_on, membership_panel, staleness
    members_on("2024-03-28")                                 # DataFrame: isin, symbol, company_name
    membership_panel(dates, columns=isins)                    # dates x ISIN booleans, the engine's `membership=`
    membership_panel(dates, index="NIFTY MIDCAP 150")         # any of INDEXES
    staleness("2026-10-24")                                   # is the last snapshot older than the last NSE review?
    python -m ros.data.nifty500                               # print the snapshot summary
"""
from __future__ import annotations

import datetime as dt
import glob
import hashlib
import json
import os
import re
from collections.abc import Iterable
from typing import Any

import pandas as pd

from ros.data import sources

INDEXES = {
    "NIFTY 500": None,                         # every row of a snapshot
    "NIFTY 50": "in_nifty50",
    "NIFTY NEXT 50": "in_next50",
    "NIFTY 100": "in_nifty100",
    "NIFTY MIDCAP 150": "in_midcap150",
    "NIFTY SMALLCAP 250": "in_smallcap250",
    # NSE defines these as unions of the above, so they follow from the same flags
    "NIFTY LARGEMIDCAP 250": ("in_nifty100", "in_midcap150"),
    "NIFTY MIDSMALLCAP 400": ("in_midcap150", "in_smallcap250"),
}
# header text in the Bifurcation sheet -> flag column (first occurrence of each header is used)
_SLEEVE_HEADERS = {"Nifty 50": "in_nifty50", "Nifty Next 50": "in_next50", "Nifty 100": "in_nifty100",
                   "Nifty Midcap 150": "in_midcap150", "Nifty Smallcap 250": "in_smallcap250"}
MIN_COMPLETE = 480          # a NIFTY 500 snapshot with fewer names is a truncated pull, not a smaller index
_ISIN_RE = re.compile(r"IN[A-Z0-9]{10}")
_UPDATE_NAME = re.compile(r"(\d{4}-\d{2}-\d{2})\.csv$", re.IGNORECASE)
SNAPSHOT_COLUMNS = ["snapshot_date", "bbg_id", "isin", "symbol", "company_name", "mapped", *_SLEEVE_HEADERS.values(),
                    "source"]


class MembershipError(ValueError):
    """The holdings workbook (or an update CSV) is not what this module documents."""


# ------------------------------------------------------------------------------------------------ reading
def _is_isin(v) -> bool:
    return isinstance(v, str) and bool(_ISIN_RE.fullmatch(v.strip()))


def _text(v) -> str | None:
    return v.strip() if isinstance(v, str) and v.strip() not in ("", "#N/A") else None


def _read_bifurcation(path: str) -> pd.DataFrame:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        # Two sheets carry the layout: "Bifurcation" (every quarter since 2012) and "N500 since Dec-2018" (a
        # shorter copy). Prefer the one named for it; otherwise the first match.
        matches = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
                cells = [c.strip() if isinstance(c, str) else c for c in row]
                if "ISIN" in cells and "Date" in cells and "Nifty Smallcap 250" in cells:
                    matches.append((ws, cells))
                    break
        if not matches:
            raise MembershipError(f"{path}: no sheet whose first six rows carry 'ISIN', 'Date' and "
                                  f"'Nifty Smallcap 250' headers (the Bifurcation sheet). Sheets: {wb.sheetnames}")
        sheet, header = next(((ws, h) for ws, h in matches if "bifurc" in ws.title.lower()), matches[0])
        col = {}
        for i, h in enumerate(header):
            if h in ("ISIN", "Date", "Ticker", "Company Name") and h not in col:
                col[h] = i
            elif h in _SLEEVE_HEADERS and _SLEEVE_HEADERS[h] not in col:
                col[_SLEEVE_HEADERS[h]] = i
            elif h in ("id().value", "Bloomberg Ticker") and "bbg" not in col:
                col["bbg"] = i
        missing = [k for k in ("ISIN", "Date", "bbg", *_SLEEVE_HEADERS.values()) if k not in col]
        if missing:
            raise MembershipError(f"{path}: Bifurcation sheet lacks columns {missing}")
        def cell(row, k):
            return row[col[k]] if k in col and col[k] < len(row) else None

        rows = []
        started = False
        for row in sheet.iter_rows(values_only=True):
            if not started:
                started = any(isinstance(c, str) and c.strip() == "ISIN" for c in row)
                continue
            d = cell(row, "Date")
            if not isinstance(d, (dt.datetime, dt.date)):
                continue
            isin = _text(cell(row, "ISIN"))
            rows.append({
                "snapshot_date": pd.Timestamp(d).normalize(), "bbg_id": _text(cell(row, "bbg")),
                "isin": isin if _is_isin(isin) else None, "symbol": _text(cell(row, "Ticker")),
                "company_name": _text(cell(row, "Company Name")),
                **{flag: _text(cell(row, flag)) is not None for flag in _SLEEVE_HEADERS.values()},
                "source": os.path.basename(path),
            })
    finally:
        wb.close()
    if not rows:
        raise MembershipError(f"{path}: the Bifurcation sheet has no dated rows")
    df = pd.DataFrame(rows)
    df["mapped"] = df["isin"].notna()
    return df[SNAPSHOT_COLUMNS]


def _read_updates(updates_dir: str) -> pd.DataFrame:
    """Extra snapshots: data/raw/nifty500_updates/YYYY-MM-DD.csv with an `isin` column."""
    frames = []
    for f in sorted(glob.glob(os.path.join(updates_dir, "*.csv"))):
        m = _UPDATE_NAME.search(os.path.basename(f))
        if not m:
            raise MembershipError(f"{f}: update files must be named YYYY-MM-DD.csv (the snapshot date)")
        d = pd.Timestamp(m.group(1))
        raw = pd.read_csv(f, dtype=str, encoding="utf-8")
        raw.columns = [c.strip().lower() for c in raw.columns]
        if "isin" not in raw.columns:
            raise MembershipError(f"{f}: needs an 'isin' column; has {list(raw.columns)}")
        raw["isin"] = raw["isin"].str.strip().str.upper()
        bad = ~raw["isin"].map(_is_isin)
        if bad.any():
            raise MembershipError(f"{f}: {int(bad.sum())} row(s) with a malformed ISIN, e.g. {raw.loc[bad, 'isin'].head(3).tolist()}")
        if raw["isin"].duplicated().any():
            raise MembershipError(f"{f}: duplicate ISINs {raw.loc[raw['isin'].duplicated(), 'isin'].head(3).tolist()}")
        if len(raw) < MIN_COMPLETE:
            raise MembershipError(f"{f}: only {len(raw)} rows; a NIFTY 500 list has about 500 (MIN_COMPLETE={MIN_COMPLETE})")
        out = pd.DataFrame({
            "snapshot_date": d, "bbg_id": None, "isin": raw["isin"],
            "symbol": raw["symbol"].str.strip() if "symbol" in raw.columns else None,
            "company_name": raw["company_name"].str.strip() if "company_name" in raw.columns else None,
            "mapped": True, "source": os.path.basename(f),
        })
        for flag in _SLEEVE_HEADERS.values():
            out[flag] = raw[flag].astype(str).str.strip().isin(["1", "true", "True", "yes"]) if flag in raw.columns else False
        frames.append(out[SNAPSHOT_COLUMNS])
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=SNAPSHOT_COLUMNS)


def _cache_key(path: str, updates_dir: str) -> str:
    h = hashlib.sha1(os.path.abspath(path).encode() + sources.file_signature(path).encode())
    for f in sorted(glob.glob(os.path.join(updates_dir, "*.csv"))):
        h.update(os.path.basename(f).encode() + sources.file_signature(f).encode())
    return h.hexdigest()[:10]


def load_snapshots(path: str | None = None, updates_dir: str | None = None, include_incomplete: bool = False,
                   out_dir: str | None = None) -> pd.DataFrame:
    """Every (snapshot_date, member) row: bbg_id, isin, symbol, company_name, mapped, in_* flags, source.

    Cached as Parquet under data/processed/nifty500/ keyed by the workbook's and the update files' size+mtime.
    """
    path = sources.require(path or sources.NIFTY500_HOLDINGS_XLSX, "NIFTY 500 holdings workbook")
    updates_dir = updates_dir or sources.NIFTY500_UPDATES_DIR
    out_dir = out_dir or os.path.join(sources.PROCESSED_DIR, "nifty500")
    os.makedirs(out_dir, exist_ok=True)
    cached = os.path.join(out_dir, f"snapshots_{_cache_key(path, updates_dir)}.parquet")
    if os.path.exists(cached):
        snaps = pd.read_parquet(cached)
    else:
        snaps = pd.concat([_read_bifurcation(path), _read_updates(updates_dir)], ignore_index=True)
        wb_last = snaps.loc[snaps["source"].str.endswith(".xlsx"), "snapshot_date"].max()
        early = snaps[(~snaps["source"].str.endswith(".xlsx")) & (snaps["snapshot_date"] <= wb_last)]
        if not early.empty:
            raise MembershipError(f"update file(s) {sorted(early['source'].unique())} are dated on or before the "
                                  f"workbook's last snapshot {wb_last.date()}; updates must be later")
        snaps = snaps.sort_values(["snapshot_date", "isin"]).reset_index(drop=True)
        snaps.to_parquet(cached, index=False)
    if not include_incomplete:
        n = snaps.groupby("snapshot_date")["isin"].size()
        keep = n[n >= MIN_COMPLETE].index
        snaps = snaps[snaps["snapshot_date"].isin(keep)].reset_index(drop=True)
    return snaps


def snapshot_summary(path: str | None = None, **kw) -> pd.DataFrame:
    """One row per snapshot: names listed, mapped to an ISIN, unmapped, complete, and each sleeve's count."""
    snaps = load_snapshots(path, include_incomplete=True, **kw)
    g = snaps.groupby("snapshot_date")
    out = pd.DataFrame({"names": g.size(), "mapped": g["mapped"].sum(), "source": g["source"].first()})
    out["unmapped"] = out["names"] - out["mapped"]
    out["complete"] = out["names"] >= MIN_COMPLETE
    for flag in _SLEEVE_HEADERS.values():
        out[flag] = g[flag].sum()
    return out


# --------------------------------------------------------------------------------------------- membership
def _flag(index: str) -> tuple[str, ...] | None:
    key = index.strip().upper()
    if key not in INDEXES:
        raise MembershipError(f"unknown index {index!r}; choose from {list(INDEXES)}")
    f = INDEXES[key]
    return None if f is None else (f,) if isinstance(f, str) else tuple(f)


def load_membership(index: str = "NIFTY 500", path: str | None = None, **kw) -> pd.DataFrame:
    """Spells: isin, symbol, company_name, start_date, end_date (exclusive; NaT = still a member at the last
    snapshot). Only ISIN-mapped names, only complete snapshots."""
    snaps = load_snapshots(path, **kw)
    flag = _flag(index)
    rows = snaps[snaps["mapped"]]
    if flag:
        rows = rows[rows[list(flag)].any(axis=1)]
    dates = sorted(snaps["snapshot_date"].unique())
    present = rows.pivot_table(index="snapshot_date", columns="isin", values="mapped", aggfunc="any") \
        .reindex(dates).fillna(False).astype(bool)
    names = rows.sort_values("snapshot_date").groupby("isin")[["symbol", "company_name"]].last()
    spells = []
    for isin in present.columns:
        s = present[isin].to_numpy()
        start = None
        for i, d in enumerate(dates):
            if s[i] and start is None:
                start = d
            elif not s[i] and start is not None:
                spells.append((isin, start, d))
                start = None
        if start is not None:
            spells.append((isin, start, pd.NaT))
    out = pd.DataFrame(spells, columns=["isin", "start_date", "end_date"])
    out["start_date"] = pd.to_datetime(out["start_date"])
    out["end_date"] = pd.to_datetime(out["end_date"])
    out = out.join(names, on="isin")
    return out[["isin", "symbol", "company_name", "start_date", "end_date"]] \
        .sort_values(["start_date", "isin"]).reset_index(drop=True)


def snapshot_dates(path: str | None = None, **kw) -> pd.DatetimeIndex:
    return pd.DatetimeIndex(sorted(load_snapshots(path, **kw)["snapshot_date"].unique()), name="snapshot_date")


def change_dates(index: str = "NIFTY 500", m: pd.DataFrame | None = None) -> pd.DatetimeIndex:
    """Every date on which the membership of `index` actually changed."""
    m = load_membership(index) if m is None else m
    d = pd.concat([m["start_date"], m["end_date"].dropna()]).drop_duplicates().sort_values()
    return pd.DatetimeIndex(d, name="date")


def _check_date(d, first: pd.Timestamp) -> pd.Timestamp:
    d = pd.Timestamp(d).normalize()
    if d < first:
        raise MembershipError(f"{d.date()} is before the first snapshot {first.date()}: no membership is known")
    return d


def members_on(d, index: str = "NIFTY 500", m: pd.DataFrame | None = None) -> pd.DataFrame:
    """Who was in `index` on day d: isin, symbol, company_name. Carried forward after the last snapshot (see
    `staleness`)."""
    m = load_membership(index) if m is None else m
    d = _check_date(d, m["start_date"].min())
    live = m[(m["start_date"] <= d) & (m["end_date"].isna() | (m["end_date"] > d))]
    return live[["isin", "symbol", "company_name"]].sort_values("isin").reset_index(drop=True)


def membership_panel(dates: Iterable, columns: Iterable[str] | None = None, key: str = "isin",
                     index: str = "NIFTY 500", m: pd.DataFrame | None = None) -> pd.DataFrame:
    """dates x securities booleans, the engine's `membership=`. `key` is 'isin' (recommended) or 'symbol'.

    A column asked for that is never a member is all False; nothing is invented for it.
    """
    if key not in ("isin", "symbol"):
        raise MembershipError("key must be 'isin' or 'symbol'")
    m = load_membership(index) if m is None else m
    idx = pd.DatetimeIndex(pd.to_datetime(list(dates))).normalize()
    if len(idx) == 0:
        return pd.DataFrame(index=idx, columns=list(columns or []), dtype=bool)
    _check_date(idx.min(), m["start_date"].min())
    cols = list(columns) if columns is not None else sorted(m[key].dropna().unique())
    out = pd.DataFrame(False, index=idx, columns=cols)
    out.index.name = "date"
    want = set(cols)
    end = m["end_date"].fillna(pd.Timestamp.max)
    for k, s, e in zip(m[key], m["start_date"], end):
        if k in want:
            out.loc[(idx >= s) & (idx < e), k] = True
    return out


def next_review_after(d) -> pd.Timestamp:
    """The first NSE semi-annual review effective date (31 March / 30 September) strictly after d."""
    d = pd.Timestamp(d).normalize()
    cands = [pd.Timestamp(d.year, 3, 31), pd.Timestamp(d.year, 9, 30), pd.Timestamp(d.year + 1, 3, 31)]
    return next(c for c in cands if c > d)


def staleness(end_date, path: str | None = None, **kw) -> dict[str, Any]:
    """Whether membership is carried forward across an NSE review the file has not seen.

    last_snapshot: the newest snapshot date. next_review: the first review effective after it. stale: True when
    `end_date` is on or after that review, i.e. the index was re-formed but the file still lists the old members.
    """
    last = snapshot_dates(path, **kw).max()
    nxt = next_review_after(last)
    end = pd.Timestamp(end_date).normalize()
    stale = end >= nxt
    msg = (f"Membership after {last.date()} is the {last.date()} list carried forward. NSE reviewed the index "
           f"effective {nxt.date()}, so holdings from {nxt.date()} to {end.date()} rest on a stale list. Drop the "
           f"post-review constituents in {sources.NIFTY500_UPDATES_DIR} as {nxt.date()}.csv (isin column)."
           if stale else
           f"Membership is current: the last snapshot {last.date()} is after the latest NSE review effective before "
           f"{end.date()}.")
    return {"last_snapshot": str(last.date()), "next_review": str(nxt.date()), "end_date": str(end.date()),
            "stale": bool(stale), "message": msg}


def provenance(path: str | None = None, **kw) -> dict[str, Any]:
    path = sources.require(path or sources.NIFTY500_HOLDINGS_XLSX, "NIFTY 500 holdings workbook")
    summ = snapshot_summary(path, **kw)
    complete = summ[summ["complete"]]
    return {
        "loader": "ros.data.nifty500", "source": path, "source_signature": sources.file_signature(path),
        "updates_dir": sources.NIFTY500_UPDATES_DIR,
        "update_files": sorted(os.path.basename(f) for f in glob.glob(os.path.join(sources.NIFTY500_UPDATES_DIR, "*.csv"))),
        "snapshots": len(summ), "complete_snapshots": len(complete),
        "skipped_incomplete": [str(d.date()) for d in summ.index[~summ["complete"]]],
        "first_snapshot": str(complete.index.min().date()), "last_snapshot": str(complete.index.max().date()),
        "unmapped_names_first": complete["unmapped"].iloc[0].item(),
        "unmapped_names_last": complete["unmapped"].iloc[-1].item(),
        "pit_status": "point_in_time",
        "read_at": sources.stamp(),
        "caveats": [
            ("Quarterly snapshots; a name is a member from the snapshot that first lists it until the first snapshot "
             "that does not. NSE's exact effective dates inside a quarter are not represented."),
            ("Members whose Bloomberg id has no ISIN in the security master (delisted/merged before the export) are "
             "left out: survivorship bias at the margin, largest in 2012-2014."),
            "Snapshots with fewer than MIN_COMPLETE names are skipped as truncated pulls.",
            "After the last snapshot the list is carried forward; see staleness().",
        ],
    }


def _cli(argv=None) -> int:
    print(json.dumps(provenance(), indent=2))
    print()
    print(snapshot_summary().to_string())
    return 0


if __name__ == "__main__":
    raise SystemExit(_cli())
