"""Ace Equity daily prices: the one stock-level price source in this repository.

THE FILE (ros/data/sources.py: ACE_PRICES_CSV)
    One row per (Accord Code, trading day) for every company Ace Equity covers, January 2000 onward, about
    9.5 million rows and 4,500 companies. Columns used:
        Accord Code                  Ace's stable numeric security id. Never reused, never renamed.
        NDP_Date                     trading day (YYYY-MM-DD)
        NDP_Open/High/Low/Close      ADJUSTED for splits, bonuses and other corporate actions (back-adjusted)
        NDP_Un Adj Open/.../Close    the raw traded prices of the day
        NDP_Mcap                     market capitalisation, Rs crore (total, not free float)
        NDP_Volume ('000)            shares traded, thousands
        NDP_Value                    traded value, Rs crore
        NDP_No of trades             number of trades
        NDP_PE, NDP_Price_BV, NDP_EV_EBIDTA, NDP_MCAP_NS  standalone valuation ratios on the day
        NDP_Consolidated ...         the same four on a consolidated basis
    Dividends are not in the file: returns from `close` are PRICE returns.

KEYS
    Every durable key in this repository is the ISIN (house rule). Ace keys by Accord Code, so the join is the
    security master in the NIFTY 500 holdings workbook (its "Sheet2": Accord Code, NSE symbol, Bloomberg code,
    ISIN, sector). `security_master()` builds that bridge and appends the Ace codes the master does not know,
    with a blank ISIN. Wide panels keyed by ISIN leave those out and report how many.

CACHE
    The CSV is 2.3 GB and takes a minute to parse, so the first call writes a Parquet copy under
    data/processed/ace/ (one per source size+mtime; a changed source file is re-parsed automatically) and every
    later call reads only the columns it needs from that. data/raw/ is never written. Delete data/processed/ace
    to force a rebuild.

CLEANING, all counted in `provenance()`
    rows whose date does not parse, and rows with a missing or non-positive adjusted close, are dropped.
    Nothing else is touched: no gap filling, no outlier clipping, no corporate-action heuristics (the adjusted
    columns already carry Ace's adjustments; `close / unadj_close` on a day is the cumulative adjustment factor).

CALENDAR
    The NSE trading calendar is the set of dates on which the file has prices (`trading_calendar()`), not
    weekdays: it includes Saturday special sessions and excludes holidays. Dates are naive calendar dates, as
    the engines expect; provenance timestamps are tz-aware Asia/Kolkata.

    from ros.data import ace
    close = ace.wide("close")                               # dates x ISIN, adjusted close
    panel = ace.panel(["close", "mcap_cr", "value_cr"])      # dict of wide frames
    long = ace.load_long(["close", "volume_000"], start="2015-01-01")
    sm = ace.security_master()                              # accord_code, isin, nse_symbol, company_name, ...
    python -m ros.data.ace                                   # build the cache, print provenance and coverage
"""
from __future__ import annotations

import hashlib
import json
import os
from collections.abc import Iterable
from typing import Any

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq

from ros.data import sources

# canonical name -> column header in the Ace export
SOURCE_COLUMNS = {
    "accord_code": "Accord Code",
    "company_name": "Company Name",
    "date": "NDP_Date",
    "open": "NDP_Open", "high": "NDP_High", "low": "NDP_Low", "close": "NDP_Close",
    "mcap_cr": "NDP_Mcap",
    "volume_000": "NDP_Volume ('000)",
    "value_cr": "NDP_Value",
    "trades": "NDP_No of trades",
    "unadj_open": "NDP_Un Adj Open", "unadj_high": "NDP_Un Adj High",
    "unadj_low": "NDP_Un Adj Low", "unadj_close": "NDP_Un Adj Close",
    "pe": "NDP_PE", "pbv": "NDP_Price_BV", "ev_ebitda": "NDP_EV_EBIDTA", "mcap_sales": "NDP_MCAP_NS",
    "pe_cons": "NDP_Consolidated PE", "pbv_cons": "NDP_Consolidated Price BV",
    "mcap_sales_cons": "NDP_Consolidated MCAP NS", "ev_ebitda_cons": "NDP_Consolidated EV EBIDTA",
}
NUMERIC_FIELDS = tuple(k for k in SOURCE_COLUMNS if k not in ("accord_code", "company_name", "date"))
_SCHEMA = pa.schema([("date", pa.timestamp("ns")), ("accord_code", pa.int32())]
                    + [(f, pa.float64()) for f in NUMERIC_FIELDS])

# security master columns in the holdings workbook's "Sheet2"
_MASTER_COLUMNS = {"Accord Code": "accord_code", "Company Name": "company_name", "CD_NSE Symbol": "nse_symbol",
                   "CD_Bloomberg Code": "bloomberg", "CD_ISIN No": "isin", "CD_Sector": "sector"}


class AceDataError(ValueError):
    """The Ace file or the security master is not what this module documents. Never downgraded to a warning."""


# ------------------------------------------------------------------------------------------------- cache
def cache_dir() -> str:
    return os.path.join(sources.PROCESSED_DIR, "ace")


def _sig(*paths: str) -> str:
    h = hashlib.sha1()
    for p in paths:
        h.update(os.path.abspath(p).encode("utf-8"))
        h.update(sources.file_signature(p).encode("ascii"))
    return h.hexdigest()[:10]


def _paths(path: str, out_dir: str) -> dict[str, str]:
    base = os.path.join(out_dir, f"prices_{_sig(path)}")
    return {"prices": base + ".parquet", "securities": base + "_securities.parquet", "prov": base + ".json"}


def _sha256(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 24), b""):
            h.update(chunk)
    return h.hexdigest()


def build_cache(path: str | None = None, out_dir: str | None = None, force: bool = False,
                chunksize: int = 1_000_000) -> dict[str, Any]:
    """Parse the CSV once into Parquet (prices, and a one-row-per-security table). Returns the provenance dict.

    Streams the file in chunks so a 2.3 GB CSV never has to fit in memory twice. Idempotent: an existing cache
    for the same source bytes (size + mtime) is reused unless `force`.
    """
    path = sources.require(path or sources.ACE_PRICES_CSV, "Ace Equity price file")
    out_dir = out_dir or cache_dir()
    os.makedirs(out_dir, exist_ok=True)
    p = _paths(path, out_dir)
    if not force and all(os.path.exists(v) for v in p.values()):
        with open(p["prov"], encoding="utf-8") as fh:
            return json.load(fh)

    header = pd.read_csv(path, nrows=0).columns.tolist()
    missing = [v for v in SOURCE_COLUMNS.values() if v not in header]
    if missing:
        raise AceDataError(f"{path}: expected Ace columns are missing: {missing}. Columns present: {header[:30]}")

    inv = {v: k for k, v in SOURCE_COLUMNS.items()}
    dropped = {"bad_date": 0, "missing_or_nonpositive_close": 0}
    sec_parts: list[pd.DataFrame] = []
    n_rows = 0
    writer = pq.ParquetWriter(p["prices"] + ".tmp", _SCHEMA, compression="snappy")
    try:
        for chunk in pd.read_csv(path, usecols=list(SOURCE_COLUMNS.values()), chunksize=chunksize,
                                 dtype={"Accord Code": "int64", "Company Name": "string"}, low_memory=False):
            chunk = chunk.rename(columns=inv)
            chunk["date"] = pd.to_datetime(chunk["date"], format="%Y-%m-%d", errors="coerce")
            bad = chunk["date"].isna()
            dropped["bad_date"] += int(bad.sum())
            chunk = chunk[~bad]
            for f in NUMERIC_FIELDS:
                chunk[f] = pd.to_numeric(chunk[f], errors="coerce").astype("float64")
            bad = chunk["close"].isna() | (chunk["close"] <= 0)
            dropped["missing_or_nonpositive_close"] += int(bad.sum())
            chunk = chunk[~bad]
            if chunk.empty:
                continue
            sec_parts.append(chunk.groupby("accord_code").agg(
                company_name=("company_name", "last"), first_date=("date", "min"),
                last_date=("date", "max"), n_days=("date", "size")).reset_index())
            out = chunk[["date", "accord_code"] + list(NUMERIC_FIELDS)].copy()
            out["accord_code"] = out["accord_code"].astype("int32")
            out["date"] = out["date"].astype("datetime64[ns]")
            writer.write_table(pa.Table.from_pandas(out, schema=_SCHEMA, preserve_index=False))
            n_rows += len(out)
    finally:
        writer.close()
    if n_rows == 0:
        os.remove(p["prices"] + ".tmp")
        raise AceDataError(f"{path}: no usable rows (every row had a bad date or a missing close)")
    os.replace(p["prices"] + ".tmp", p["prices"])

    sec = pd.concat(sec_parts, ignore_index=True)
    sec = sec.sort_values(["accord_code", "last_date"]).groupby("accord_code").agg(
        company_name=("company_name", "last"), first_date=("first_date", "min"),
        last_date=("last_date", "max"), n_days=("n_days", "sum")).reset_index()
    sec["accord_code"] = sec["accord_code"].astype("int32")
    sec.to_parquet(p["securities"], index=False)

    dates = pq.read_table(p["prices"], columns=["date"]).column("date").to_pandas()
    prov = {
        "loader": "ros.data.ace", "source": path, "source_sha256": _sha256(path),
        "source_size_bytes": os.path.getsize(path), "source_signature": sources.file_signature(path),
        "cache": p["prices"], "built_at": sources.stamp(),
        "rows": int(n_rows), "securities": len(sec),
        "first_date": str(dates.min().date()), "last_date": str(dates.max().date()),
        "trading_days": int(dates.nunique()), "dropped_rows": dropped,
        "fields": list(NUMERIC_FIELDS),
        "notes": [
            "open/high/low/close are Ace's back-adjusted prices; unadj_* are the raw traded prices.",
            "Price returns only: no dividends in the file.",
            "mcap_cr is total market capitalisation in Rs crore, not free float.",
            "Keyed by Accord Code; ISIN comes from the NIFTY 500 holdings workbook's security master.",
        ],
    }
    with open(p["prov"], "w", encoding="utf-8") as fh:
        json.dump(prov, fh, indent=2)
    return prov


def provenance(path: str | None = None, out_dir: str | None = None) -> dict[str, Any]:
    return build_cache(path, out_dir)


def _cache_paths(path: str | None, out_dir: str | None) -> dict[str, str]:
    build_cache(path, out_dir)
    return _paths(path or sources.ACE_PRICES_CSV, out_dir or cache_dir())


# ---------------------------------------------------------------------------------------- security master
def _read_sheet2(holdings_path: str) -> pd.DataFrame:
    import openpyxl

    wb = openpyxl.load_workbook(holdings_path, read_only=True, data_only=True)
    try:
        sheet = None
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
                if "Accord Code" in row and "CD_ISIN No" in row:
                    sheet, header = ws, list(row)
                    break
            if sheet is not None:
                break
        if sheet is None:
            raise AceDataError(f"{holdings_path}: no sheet with 'Accord Code' and 'CD_ISIN No' headers in its "
                               f"first six rows (the security master). Sheets: {wb.sheetnames}")
        idx = {_MASTER_COLUMNS[h]: i for i, h in enumerate(header) if h in _MASTER_COLUMNS}
        missing = set(_MASTER_COLUMNS.values()) - set(idx)
        if missing:
            raise AceDataError(f"{holdings_path}: security master lacks columns {sorted(missing)}")
        rows = []
        started = False
        for row in sheet.iter_rows(values_only=True):
            if not started:
                started = "Accord Code" in row
                continue
            code = row[idx["accord_code"]] if idx["accord_code"] < len(row) else None
            if code is None or isinstance(code, str) and not code.strip().isdigit():
                continue
            rows.append({k: (row[i] if i < len(row) else None) for k, i in idx.items()})
    finally:
        wb.close()
    df = pd.DataFrame(rows, columns=list(_MASTER_COLUMNS.values()))
    df["accord_code"] = pd.to_numeric(df["accord_code"], errors="coerce").astype("Int64")
    df = df.dropna(subset=["accord_code"])
    for c in ("company_name", "nse_symbol", "bloomberg", "isin", "sector"):
        df[c] = df[c].map(lambda v: v.strip() if isinstance(v, str) and v.strip() not in ("", "#N/A") else None)
    bad_isin = df["isin"].notna() & ~df["isin"].str.fullmatch(r"IN[A-Z0-9]{10}")
    df.loc[bad_isin, "isin"] = None
    dup = df["accord_code"].duplicated(keep=False)
    if dup.any():
        raise AceDataError(f"security master: Accord Code repeated: {sorted(df.loc[dup, 'accord_code'].unique())[:10]}")
    dup = df["isin"].notna() & df["isin"].duplicated(keep=False)
    if dup.any():
        raise AceDataError(f"security master: ISIN shared by several Accord Codes: "
                           f"{sorted(df.loc[dup, 'isin'].unique())[:10]}")
    return df


def security_master(holdings_path: str | None = None, prices_path: str | None = None,
                    out_dir: str | None = None) -> pd.DataFrame:
    """Accord Code <-> ISIN <-> NSE symbol, for every security in the master OR in the Ace file.

    Columns: accord_code, isin, nse_symbol, bloomberg, company_name, sector, in_ace, first_date, last_date,
    n_days. `isin` is blank for Ace codes the master does not know (mostly companies delisted before the
    master was exported); those cannot be joined to index membership and are excluded from ISIN-keyed panels.
    Written to data/processed/ace/security_master.csv for inspection.
    """
    holdings_path = sources.require(holdings_path or sources.NIFTY500_HOLDINGS_XLSX, "NIFTY 500 holdings workbook")
    out_dir = out_dir or cache_dir()
    p = _cache_paths(prices_path, out_dir)
    cached = os.path.join(out_dir, f"security_master_{_sig(holdings_path, prices_path or sources.ACE_PRICES_CSV)}.parquet")
    if os.path.exists(cached):
        return pd.read_parquet(cached)
    master = _read_sheet2(holdings_path)
    ace = pd.read_parquet(p["securities"])
    ace["accord_code"] = ace["accord_code"].astype("Int64")
    sm = master.merge(ace.rename(columns={"company_name": "ace_name"}), on="accord_code", how="outer")
    sm["in_ace"] = sm["n_days"].notna()
    sm["company_name"] = sm["company_name"].fillna(sm["ace_name"])
    sm = sm.drop(columns=["ace_name"]).sort_values("accord_code").reset_index(drop=True)
    sm["accord_code"] = sm["accord_code"].astype("int64")
    sm.to_parquet(cached, index=False)
    sm.to_csv(os.path.join(out_dir, "security_master.csv"), index=False)
    return sm


# ------------------------------------------------------------------------------------------------ loading
def load_long(fields: Iterable[str] = ("close",), path: str | None = None, out_dir: str | None = None,
              start=None, end=None, codes: Iterable[int] | None = None, with_isin: bool = True) -> pd.DataFrame:
    """Long frame: date, accord_code[, isin], *fields. Reads only those columns from the Parquet cache."""
    fields = [fields] if isinstance(fields, str) else list(fields)
    unknown = [f for f in fields if f not in NUMERIC_FIELDS]
    if unknown:
        raise AceDataError(f"unknown field(s) {unknown}; choose from {list(NUMERIC_FIELDS)}")
    p = _cache_paths(path, out_dir)
    filters = []
    if start is not None:
        filters.append(("date", ">=", pd.Timestamp(start).to_datetime64()))
    if end is not None:
        filters.append(("date", "<=", pd.Timestamp(end).to_datetime64()))
    if codes is not None:
        filters.append(("accord_code", "in", [int(c) for c in codes]))
    df = pq.read_table(p["prices"], columns=["date", "accord_code"] + fields,
                       filters=filters or None).to_pandas()
    if with_isin:
        sm = security_master(prices_path=path, out_dir=out_dir)[["accord_code", "isin"]]
        sm["accord_code"] = sm["accord_code"].astype("int32")
        df = df.merge(sm, on="accord_code", how="left")
        df = df[["date", "accord_code", "isin"] + fields]
    return df.sort_values(["date", "accord_code"]).reset_index(drop=True)


def wide(field: str = "close", key: str = "isin", path: str | None = None, out_dir: str | None = None,
         start=None, end=None) -> pd.DataFrame:
    """Dates x security panel of one field. `key` is 'isin' (default; unmapped codes dropped) or 'accord_code'."""
    if key not in ("isin", "accord_code"):
        raise AceDataError("key must be 'isin' or 'accord_code'")
    df = load_long([field], path, out_dir, start, end, with_isin=(key == "isin"))
    if key == "isin":
        df = df.dropna(subset=["isin"])
    out = df.pivot(index="date", columns=key, values=field).sort_index()
    out.index.name = "date"
    out.columns.name = key
    return out


def panel(fields: Iterable[str] = ("close",), key: str = "isin", **kw) -> dict[str, pd.DataFrame]:
    """Several wide panels on one aligned (dates x securities) grid."""
    fields = [fields] if isinstance(fields, str) else list(fields)
    out = {f: wide(f, key=key, **kw) for f in fields}
    idx = out[fields[0]].index
    cols = out[fields[0]].columns
    for f in fields[1:]:
        idx, cols = idx.union(out[f].index), cols.union(out[f].columns)
    return {f: v.reindex(index=idx, columns=cols) for f, v in out.items()}


def trading_calendar(path: str | None = None, out_dir: str | None = None,
                     min_securities: int = 100) -> pd.DatetimeIndex:
    """Every date with at least `min_securities` priced names: the NSE calendar as the file saw it."""
    p = _cache_paths(path, out_dir)
    d = pq.read_table(p["prices"], columns=["date"]).column("date").to_pandas()
    counts = d.value_counts()
    return pd.DatetimeIndex(sorted(counts[counts >= min_securities].index), name="date")


def unmapped_report(path: str | None = None, out_dir: str | None = None) -> pd.DataFrame:
    """Ace securities with no ISIN in the master: who they are and when they last traded."""
    sm = security_master(prices_path=path, out_dir=out_dir)
    return sm[sm["in_ace"] & sm["isin"].isna()][["accord_code", "company_name", "first_date", "last_date", "n_days"]] \
        .sort_values("last_date", ascending=False).reset_index(drop=True)


def coverage_report(path: str | None = None, out_dir: str | None = None) -> pd.DataFrame:
    """Per calendar year: trading days, securities priced, and how many of those carry an ISIN."""
    df = load_long(["close"], path, out_dir)
    df["year"] = df["date"].dt.year
    g = df.groupby("year")
    out = pd.DataFrame({
        "trading_days": g["date"].nunique(),
        "securities": g["accord_code"].nunique(),
        "with_isin": df[df["isin"].notna()].groupby("year")["accord_code"].nunique(),
    }).fillna(0).astype(int)
    out["isin_share"] = (out["with_isin"] / out["securities"]).round(3)
    return out


def _cli(argv=None) -> int:
    import argparse

    ap = argparse.ArgumentParser(description="Build the Ace price cache and describe it.")
    ap.add_argument("--force", action="store_true", help="re-parse the CSV even if a cache exists")
    args = ap.parse_args(argv)
    prov = build_cache(force=args.force)
    print(json.dumps({k: v for k, v in prov.items() if k != "notes"}, indent=2))
    sm = security_master()
    print(f"\nsecurity master: {len(sm)} securities, {int(sm['isin'].notna().sum())} with ISIN, "
          f"{int((sm['in_ace'] & sm['isin'].isna()).sum())} Ace codes without ISIN")
    print("\ncoverage by year:")
    print(coverage_report().to_string())
    return 0


if __name__ == "__main__":
    raise SystemExit(_cli())
