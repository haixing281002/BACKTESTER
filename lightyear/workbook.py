"""The results workbook for any Lightyear run, built to the xlsx skill's rules from the run's hand-off files.

Sheets
  Summary        metrics as formulas over named ranges; Gate B with live PASS/FAIL formulas against editable
                 (yellow) thresholds; colour key; two Excel charts
  Returns        start row (100 at the close before the first return), daily returns (blue), value and
                 drawdown formulas (black)
  Comparators    daily returns of the simpler alternatives (long leg alone, equal-weight list, the index...),
                 each with value/drawdown formulas and a formula metrics table          [if comparators.csv]
  Rebalances     one row per rebalance date: names long/short, gross long/short, net, as formulas over Holdings
  Holdings       every position at every rebalance (target weights)                     [if holdings.csv]
  Trades         every change between rebalances: BUY / SELL / SHORT / COVER / ADD / TRIM, with the weight change
                 as a formula                                                           [if trades or holdings]
  Positives & negatives   measured from the returns, from the run's tests, and Claude's analyst read
  Validation, one sheet per table CSV the run listed, Notes
Colour code from universal_backtester/xlsx_style.py. Recalculate afterwards with scripts/xlsx_recalc_libreoffice.py.
"""
import os
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, PatternFill

from universal_backtester.xlsx_style import (DATE, F_BASE, F_BOLD, F_INPUT, F_LINK, F_NOTE, F_SUB, F_TITLE,
                                             FILL_FAIL, FILL_KEY, FILL_PASS, INT, NUM1, NUM2, PCT, PCT2, drawdown_scale,
                                             header, legend, name, pass_fail, put, sign_colour, widths)

from . import gates
from .charts import start_date
from .paths import REPO

C_STRAT, C_BENCH = "1F4E79", "D98324"
FMT = {"pct": PCT2, "num": NUM2, "int": INT, "text": None}
PCT_NAME = re.compile(r"cagr|vol|drawdown|\bdd\b|return|\bret\b|alpha|turnover|weight|%|pct", re.I)
FILL_ADV = PatternFill("solid", bgColor="FFEB9C", fgColor="FFEB9C")
RET_FMT = "0.000%;(0.000%);-"


def _num_or_text(v):
    if isinstance(v, bool):
        return str(v)
    try:
        return float(v)
    except (TypeError, ValueError):
        return "" if v is None else str(v)


def _sheet_title(title, used):
    t = re.sub(r"[\[\]\:\*\?\/\\]", " ", str(title)).strip()[:28] or "Table"
    base, k = t, 2
    while t.lower() in used:
        t = f"{base[:25]} {k}"
        k += 1
    used.add(t.lower())
    return t


def _line_chart(title, ws_data, col_s, col_b, last, anchor, ws_target, y_fmt, y_title):
    ch = LineChart()
    ch.title, ch.height, ch.width = title, 8.5, 17
    ch.y_axis.title, ch.y_axis.number_format = y_title, y_fmt
    ch.x_axis.number_format = "mmm yy"
    ch.x_axis.tickLblSkip = max(1, (last - 1) // 8)
    ch.x_axis.delete = ch.y_axis.delete = False
    for col in (col_s, col_b):
        ch.add_data(Reference(ws_data, min_col=col, min_row=1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws_data, min_col=1, min_row=2, max_row=last))
    for s, colour in zip(ch.series, (C_STRAT, C_BENCH)):
        s.graphicalProperties.line.solidFill = colour
        s.graphicalProperties.line.width = 19000
        s.smooth = False
    ch.legend.position = "b"
    ws_target.add_chart(ch, anchor)


def _cost_basis(results):
    """How the strategy column was costed, read from results.json (never assumed)."""
    bp = results.get("headline_cost_bp")
    if bp is None:
        return "strategy as written by the run (cost basis not stated)"
    return "strategy gross (no trading costs)" if float(bp) == 0 else f"strategy after {float(bp):g}bp round-trip costs"


MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
TRADE_HEADERS = {"isin": "ISIN", "signal_date": "Signal date (ranking close)", "execution_date": "Execution date",
                 "price_date": "Price date (bhavcopy used)", "price": "Price used (Rs)", "price_basis": "Price source",
                 "shares": "Shares traded", "value_rs": "Trade value (Rs)", "why": "Why"}
DECISION_FILL = {"BOUGHT": FILL_PASS, "KEPT": PatternFill("solid", fgColor="DDEBF7"), "SOLD": FILL_FAIL,
                 "SKIPPED": PatternFill("solid", fgColor="FFEB9C"), "NOT NEEDED": PatternFill("solid", fgColor="E7E6E6")}


def _monthly_sheet(wb, df, s_name, b_name, has_trades):
    """Month-by-month returns as formulas over the Returns sheet, plus calendar grids."""
    ws = wb.create_sheet("Monthly returns")
    put(ws, "A1", "Month-by-month returns", F_TITLE)
    put(ws, "A2", "Each month runs from the previous month's last close to this month's last close: the return of the book "
        "formed at the previous month-end close. The first month is partial (from the start date). Month-end dates are "
        "inputs (blue); every return, value and count is a formula over the Returns and Trades sheets.", F_NOTE)
    me = df.groupby([df["date"].dt.year, df["date"].dt.month])["date"].max().sort_values().tolist()
    cols = ["Month", "Month-end close", "Book held from the close of", f"{s_name} return", f"{b_name} return",
            "Difference", f"{s_name} value (100 = start)", f"{b_name} value (100 = start)"]
    if has_trades:
        cols += ["Names bought at that close", "Names sold at that close"]
    cols += ["Year", "Month no."]
    header(ws, 4, cols, height=44)
    f0, f1 = 5, 4 + len(me)
    yc, mc = ("K", "L") if has_trades else ("I", "J")
    for k, d in enumerate(me):
        i = f0 + k
        put(ws, f"A{i}", f"{MONTHS[d.month - 1]} {d.year}", F_BASE, border=True)
        put(ws, f"B{i}", d.to_pydatetime(), F_INPUT, DATE, border=True)
        put(ws, f"C{i}", "=Start_Date" if k == 0 else f"=B{i-1}", F_LINK if k == 0 else F_BASE, DATE, border=True)
        put(ws, f"G{i}", f"=INDEX(Val_Strat,MATCH(B{i},Dates,0))", F_LINK, "0.00", border=True)
        put(ws, f"H{i}", f"=INDEX(Val_Bench,MATCH(B{i},Dates,0))", F_LINK, "0.00", border=True)
        put(ws, f"D{i}", f"=G{i}/100-1" if k == 0 else f"=G{i}/G{i-1}-1", F_BASE, PCT, border=True)
        put(ws, f"E{i}", f"=H{i}/100-1" if k == 0 else f"=H{i}/H{i-1}-1", F_BASE, PCT, border=True)
        put(ws, f"F{i}", f"=D{i}-E{i}", F_BASE, PCT, border=True)
        if has_trades:
            put(ws, f"I{i}", f'=COUNTIFS(T_Date,C{i},T_Action,"BUY")', F_LINK, INT, border=True)
            put(ws, f"J{i}", f'=COUNTIFS(T_Date,C{i},T_Action,"SELL")', F_LINK, INT, border=True)
        put(ws, f"{yc}{i}", f"=YEAR(B{i})", F_BASE, "0", border=True)
        put(ws, f"{mc}{i}", f"=MONTH(B{i})", F_BASE, "0", border=True)
    sign_colour(ws, f"D{f0}:F{f1}")
    r = f1 + 2
    put(ws, f"A{r}", "Over all months", F_SUB)
    stats = [("Months", f"=COUNT(D{f0}:D{f1})", INT), (f"Months {s_name} beat {b_name}", f'=COUNTIF(F{f0}:F{f1},">0")', INT),
             ("Share of months it beat the index", f"=B{r+2}/B{r+1}", PCT),
             (f"Best month ({s_name})", f"=MAX(D{f0}:D{f1})", PCT), (f"Worst month ({s_name})", f"=MIN(D{f0}:D{f1})", PCT),
             (f"Average month ({s_name})", f"=AVERAGE(D{f0}:D{f1})", PCT),
             (f"Average month ({b_name})", f"=AVERAGE(E{f0}:E{f1})", PCT)]
    for k, (lab, fml, fmt) in enumerate(stats, start=1):
        put(ws, f"A{r+k}", lab, F_BASE, border=True)
        put(ws, f"B{r+k}", fml, F_BASE, fmt, border=True)

    years = sorted({d.year for d in me})
    gc = 15                                       # calendar grids start in column O
    L = lambda c: ws.cell(row=1, column=c).column_letter
    top = 4
    for title, src, val in ((f"{s_name}: calendar view", "D", "G"), (f"{b_name}: calendar view", "E", "H"),
                            ("Difference (strategy minus index)", "F", None)):
        put(ws, f"{L(gc)}{top-1}", title, F_SUB)
        header(ws, top, ["Year"] + MONTHS + ["Year"], col=gc)
        for k, y in enumerate(years):
            i = top + 1 + k
            put(ws, f"{L(gc)}{i}", y, F_INPUT, "0", border=True)
            for m in range(1, 13):
                cell = f"{L(gc + m)}{i}"
                put(ws, cell, f'=IF(COUNTIFS(${yc}${f0}:${yc}${f1},${L(gc)}{i},${mc}${f0}:${mc}${f1},{m})=0,"",'
                              f'SUMIFS(${src}${f0}:${src}${f1},${yc}${f0}:${yc}${f1},${L(gc)}{i},${mc}${f0}:${mc}${f1},{m}))',
                    F_BASE, PCT, border=True)
            yr = f"{L(gc + 13)}{i}"
            if val:
                put(ws, yr, f"=INDEX(${val}${f0}:${val}${f1},MATCH(${L(gc)}{i},${yc}${f0}:${yc}${f1},1))/"
                            f"IFERROR(INDEX(${val}${f0}:${val}${f1},MATCH(${L(gc)}{i}-1,${yc}${f0}:${yc}${f1},1)),100)-1",
                    F_BOLD, PCT, border=True)
            else:
                put(ws, yr, f"={L(gc + 13)}{top + 1 + k - 2 * (len(years) + 3)}-{L(gc + 13)}{top + 1 + k - (len(years) + 3)}",
                    F_BOLD, PCT, border=True)
        sign_colour(ws, f"{L(gc + 1)}{top + 1}:{L(gc + 13)}{top + len(years)}")
        top += len(years) + 3
    put(ws, f"{L(gc)}{top}", "Year column = the whole calendar year (first and last years are part-years).", F_NOTE)
    ws.freeze_panes = "B5"
    widths(ws, [11, 14, 16, 14, 14, 12, 14, 14] + ([12, 12] if has_trades else []) + [8, 8, 3, 3] + [8] + [8] * 13)
    return ws


def _selection_sheet(wb, sel, n_hold=20):
    """'Why these stocks': one block per rebalance, from selection.csv written by the run."""
    s = sel.copy()
    s["signal_date"] = pd.to_datetime(s["signal_date"])
    ws = wb.create_sheet("Why these 20")
    put(ws, "A1", "Why these stocks: every rebalance explained", F_TITLE)
    put(ws, "A2", "At each month-end close every stock on that day's NIFTY 500 list with 120 days of prices is scored: "
        "annualised 120-day log-price slope x R-squared, so a steady climb scores higher than a jumpy one. A stock is ruled "
        "out if it closes below its 100-day average or had a one-day move over the limit (up or down) in the last 120 days. "
        "Going down the ranking, the first 20 that pass are held at 5% each.", F_NOTE)
    put(ws, "A3", "BOUGHT = new pick, KEPT = carried over and reset to 5%, SOLD = left the book (with the reason), "
        "SKIPPED = ranked high enough but failed a filter, NOT NEEDED = eligible just below the 20th pick (first 5 shown). "
        "Rank = place among all scored list members before the filters. Close = raw NSE bhavcopy CLOSE on the "
        "execution date, the price every trade used. All values are inputs from selection.csv written by the run; the "
        "Monthly ledger sheet has the full ~50-name list per month.", F_NOTE)
    for r in (2, 3):
        ws.row_dimensions[r].height = 42
        ws[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:N2"); ws.merge_cells("A3:N3")
    dates = sorted(s["signal_date"].unique())
    put(ws, "A5", "Jump to a rebalance (click):", F_SUB)
    idx_rows = {}
    first_block = 7 + (len(dates) + 5) // 6
    cols = ["Decision", "Pick no.", "Rank", "Symbol", "Score", "Slope (a year)", "R-squared", "Close vs 100-day avg",
            "Biggest 1-day move (120 days)", "Date of that move", "Close used (Rs, raw)", "Weight before", "Weight after", "Why"]
    row = first_block
    for d in dates:
        b = s[s["signal_date"] == d]
        cut = b["rank_of_20th_pick"].dropna()
        cut = int(cut.iloc[0]) if len(cut) else None
        nn = b[b["decision"] == "NOT NEEDED"].sort_values("rank").head(5)
        sk = b[(b["decision"] == "SKIPPED") & ((b["rank"] <= cut) if cut else True)]
        keep = pd.concat([b[b["decision"].isin(["BOUGHT", "KEPT"])].sort_values("pick_no"),
                          b[b["decision"].isin(["SOLD", "Not scored"])].sort_values("rank"), sk.sort_values("rank"), nn])
        nb, nk, ns = (b["decision"] == "BOUGHT").sum(), (b["decision"] == "KEPT").sum(), (b["decision"] == "SOLD").sum()
        ex = b["execution_date"].iloc[0] if "execution_date" in b else d.date()
        idx_rows[d] = row
        put(ws, f"A{row}", f"Rebalance at the close of {pd.Timestamp(d).date()} (trades priced at the {ex} close). "
            f"Bought {nb}, kept {nk}, sold {ns}" + (f"; the 20th pick was rank {cut}." if cut else "."),
            F_SUB, fill=PatternFill("solid", fgColor="DDEBF7"))
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(cols))
        header(ws, row + 1, cols, height=32)
        for k, x in enumerate(keep.itertuples(index=False), start=row + 2):
            g = lambda c, dflt=None: getattr(x, c, dflt)
            dec = str(g("decision", ""))
            put(ws, f"A{k}", dec, F_BOLD, fill=DECISION_FILL.get(dec), border=True)
            vals = [(g("pick_no"), INT), (g("rank"), INT), (g("symbol"), None), (g("score"), NUM2), (g("annualised_slope"), PCT),
                    (g("r_squared"), NUM2), (g("pct_vs_100d_avg"), PCT), (g("biggest_move_120d"), PCT),
                    (g("biggest_move_date"), None), (g("close_raw"), NUM2), (g("weight_before"), PCT), (g("weight_after"), PCT)]
            for j, (v, fmt) in enumerate(vals, start=2):
                v = _num_or_text(v)
                if isinstance(v, float) and v != v:
                    v = ""
                put(ws, f"{ws.cell(row=1, column=j).column_letter}{k}", v, F_INPUT, fmt if isinstance(v, float) else None,
                    border=True)
            put(ws, f"N{k}", str(g("why", "")), F_BASE, border=True, align=Alignment(wrap_text=True, vertical="top"))
        row = row + 2 + len(keep) + 2
    for n, d in enumerate(dates):
        rr, cc = 6 + n // 6, 1 + 2 * (n % 6)
        c = ws.cell(row=rr, column=cc, value=str(pd.Timestamp(d).date()))
        c.hyperlink = f"#'Why these 20'!A{idx_rows[d]}"
        c.font = F_LINK
    ws.freeze_panes = ws.cell(row=first_block, column=1)
    widths(ws, [12, 9, 7, 15, 9, 12, 10, 13, 15, 13, 13, 12, 12, 95])
    return ws


def derive_trades(h):
    """Trades implied by target weights at consecutive rebalances (the book drifts with prices in between)."""
    h = h.copy()
    h["date"] = pd.to_datetime(h["date"])
    piv = h.pivot_table(index="date", columns="symbol", values="weight", aggfunc="sum").fillna(0.0).sort_index()
    rows, prev = [], pd.Series(0.0, index=piv.columns)
    for dte, cur in piv.iterrows():
        for sym in piv.columns:
            b, a = float(prev[sym]), float(cur[sym])
            if abs(a - b) < 1e-9:
                continue
            if b == 0:
                act = "BUY" if a > 0 else "SHORT"
            elif a == 0:
                act = "SELL" if b > 0 else "COVER"
            elif (a > 0) != (b > 0):
                act = "FLIP"
            else:
                act = "ADD" if abs(a) > abs(b) else "TRIM"
            rows.append({"date": dte, "symbol": sym, "action": act, "weight_before": b, "weight_after": a})
        prev = cur
    return pd.DataFrame(rows, columns=["date", "symbol", "action", "weight_before", "weight_after"])


def build(run_dir, results, df, title, out_path, hurdle=0.06, holdings=None, trades=None, comparators=None,
          observations=None, review=None, selection=None):
    s_name = results.get("strategy_name") or "Strategy"
    b_name = results.get("benchmark_name") or "NIFTY 500"
    df = df.sort_values("date").reset_index(drop=True)
    n = len(df)
    first, last = 3, n + 2           # row 2 is the start row; data rows 3..n+2
    wb = Workbook()

    # ---------------------------------------------------------------- Returns
    wr = wb.active
    wr.title = "Returns"
    header(wr, 1, ["Date", f"{s_name} daily return", f"{b_name} daily return", f"{s_name} value (100 = start)",
                   f"{b_name} value (100 = start)", f"{s_name} drawdown", f"{b_name} drawdown"], height=44)
    wr.cell(row=2, column=1, value=start_date(df["date"]).to_pydatetime()).number_format = DATE
    wr.cell(row=2, column=1).font = F_INPUT
    for col in (4, 5):
        c = wr.cell(row=2, column=col, value=100)
        c.font, c.number_format = F_INPUT, "0.00"
    for col in (6, 7):
        c = wr.cell(row=2, column=col, value=0)
        c.font, c.number_format = F_BASE, PCT
    put(wr, "B2", "start: value 100 at the close before the first return", F_NOTE)
    for i, row in enumerate(df.itertuples(index=False), start=first):
        wr.cell(row=i, column=1, value=row.date.to_pydatetime()).number_format = DATE
        for col, v in ((2, row.strategy), (3, row.benchmark)):
            c = wr.cell(row=i, column=col, value=0.0 if pd.isna(v) else float(v))
            c.font, c.number_format = F_INPUT, RET_FMT
        wr.cell(row=i, column=4, value=f"=D{i-1}*(1+B{i})")
        wr.cell(row=i, column=5, value=f"=E{i-1}*(1+C{i})")
        wr.cell(row=i, column=6, value=f"=D{i}/MAX(D$2:D{i})-1")
        wr.cell(row=i, column=7, value=f"=E{i}/MAX(E$2:E{i})-1")
        for col in (1, 4, 5, 6, 7):
            wr.cell(row=i, column=col).font = F_BASE
        wr.cell(row=i, column=4).number_format = wr.cell(row=i, column=5).number_format = "0.00"
        wr.cell(row=i, column=6).number_format = wr.cell(row=i, column=7).number_format = PCT
    wr.freeze_panes = "B3"
    widths(wr, [12, 22, 22, 24, 24, 22, 22])
    drawdown_scale(wr, f"F3:G{last}")
    put(wr, f"A{last + 2}", f"Source: {os.path.relpath(os.path.join(run_dir, 'daily_returns.csv'), REPO)} (written by the run). "
        f"Blue = hardcoded daily returns, {_cost_basis(results)}. Black = formulas. Row 2 is the starting point: 100 at the "
        "close one business day before the first return, so every return counts.", F_NOTE)
    name(wb, "Dates", f"Returns!$A$2:$A${last}")
    name(wb, "Start_Date", "Returns!$A$2")
    for label, col in (("Ret_Strat", "B"), ("Ret_Bench", "C")):
        name(wb, label, f"Returns!${col}${first}:${col}${last}")
    for label, col in (("Val_Strat", "D"), ("Val_Bench", "E"), ("DD_Strat", "F"), ("DD_Bench", "G")):
        name(wb, label, f"Returns!${col}$2:${col}${last}")

    # ---------------------------------------------------------------- Summary
    ws = wb.create_sheet("Summary", 0)
    put(ws, "A1", title, F_TITLE)
    put(ws, "A2", results.get("headline", ""), F_NOTE)
    put(ws, "A4", "Hurdle rate for Sharpe", F_BOLD)
    put(ws, "B4", hurdle, F_INPUT, PCT, FILL_KEY, border=True,
        comment="Key assumption: flat 6% proxy for the Indian short rate (the repo's convention).")
    put(ws, "C4", "Flat 6% proxy, the repo's convention. Change it and the Sharpe row updates.", F_NOTE)
    name(wb, "Hurdle", "Summary!$B$4")
    header(ws, 6, ["Metric", f"Strategy: {s_name}", f"Benchmark: {b_name}", f"Strategy minus {b_name}"])
    mrows = [
        ("Start (value = 100)", "=Start_Date", "=Start_Date", DATE, False, True),
        ("End date", "=MAX(Dates)", "=MAX(Dates)", DATE, False, True),
        ("Years", "=(B8-B7)/365.25", "=(C8-C7)/365.25", NUM2, False, False),
        ("Trading days", "=COUNT(Ret_Strat)", "=COUNT(Ret_Bench)", INT, False, True),
        ("Final value (100 = start)", "=INDEX(Val_Strat,ROWS(Val_Strat))", "=INDEX(Val_Bench,ROWS(Val_Bench))", NUM1, True, True),
        ("CAGR", "=(B11/100)^(1/B9)-1", "=(C11/100)^(1/C9)-1", PCT, True, False),
        ("Annualised volatility", "=STDEV(Ret_Strat)*SQRT(252)", "=STDEV(Ret_Bench)*SQRT(252)", PCT, True, True),
        ("Sharpe vs hurdle", "=(B12-Hurdle)/B13", "=(C12-Hurdle)/C13", NUM2, True, False),
        ("Maximum drawdown", "=-MIN(DD_Strat)", "=-MIN(DD_Bench)", PCT, True, True),
    ]
    for i, (label, fs, fb, fmt, ex, link) in enumerate(mrows, start=7):
        put(ws, f"A{i}", label, F_BASE, border=True)
        put(ws, f"B{i}", fs, F_LINK if link else F_BASE, fmt, border=True)
        put(ws, f"C{i}", fb, F_LINK if link else F_BASE, fmt, border=True)
        put(ws, f"D{i}", f"=B{i}-C{i}" if ex else None, F_BASE, fmt, border=True)
    for r in (11, 12, 14):
        sign_colour(ws, f"D{r}")
    for r in (13, 15):   # more volatility / deeper drawdown than the benchmark is bad
        ws.conditional_formatting.add(f"D{r}", CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_FAIL))
        ws.conditional_formatting.add(f"D{r}", CellIsRule(operator="lessThan", formula=["0"], fill=FILL_PASS))
    put(ws, "A16", "CAGR on calendar years from the start row (days / 365.25); volatility = stdev of daily returns x "
        "sqrt(252). Rows 7 to 15 are formulas over the Returns sheet (green = pulls from another sheet).", F_NOTE)

    header(ws, 18, ["Gate B criterion (ros/governance/gates.py)", "Result", "Value", "Rule", "Threshold (editable)",
                    "Counts toward decision"])
    grows = gates.rows(results["gate_b"].get("criteria", []))
    for i, g in enumerate(grows, start=19):
        put(ws, f"A{i}", g["name"], F_BASE, border=True)
        if g["value"] is not None and g["threshold"] is not None and g["op"]:
            fmt = PCT if "%" in str(g["threshold_raw"]) or "%" in str(g["value_raw"]) else NUM2
            put(ws, f"C{i}", g["value"], F_INPUT, fmt, border=True, comment=g.get("evidence") or None)
            put(ws, f"D{i}", ("|value| " if g["abs"] else "value ") + g["op"], F_INPUT, border=True)
            put(ws, f"E{i}", g["threshold"], F_INPUT, fmt, FILL_KEY, border=True)
            lhs = f"ABS(C{i})" if g["abs"] else f"C{i}"
            put(ws, f"B{i}", f'=IF({lhs}{g["op"]}E{i},"PASS",IF(F{i}="advisory","FAIL (advisory)","FAIL"))',
                F_BOLD, border=True, align=Alignment(horizontal="center"))
        else:
            put(ws, f"C{i}", str(g["value_raw"]), F_INPUT, border=True, comment=g.get("evidence") or None)
            put(ws, f"D{i}", "", F_INPUT, border=True)
            put(ws, f"E{i}", str(g["threshold_raw"]), F_INPUT, border=True)
            put(ws, f"B{i}", g["result"], F_BOLD, border=True, align=Alignment(horizontal="center"),
                comment="Not computed: there was nothing to compare against in this run.")
        put(ws, f"F{i}", "blocking" if g["blocking"] else "advisory", F_INPUT, border=True)
    g_last = 18 + len(grows)
    if grows:
        pass_fail(ws, f"B19:B{g_last}")
        ws.conditional_formatting.add(f"B19:B{g_last}", CellIsRule(operator="equal", formula=['"FAIL (advisory)"'], fill=FILL_ADV))
    r = g_last + 1
    put(ws, f"A{r}", "Blocking criteria passed", F_BOLD, border=True)
    put(ws, f"B{r}", f'=COUNTIFS(B19:B{g_last},"PASS",F19:F{g_last},"blocking")&" of "&COUNTIF(F19:F{g_last},"blocking")',
        F_BOLD, border=True, align=Alignment(horizontal="center"))
    put(ws, f"A{r+1}", "Evidence supports", F_BOLD, border=True)
    put(ws, f"B{r+1}", results.get("evidence_supports", ""), F_BOLD, border=True, align=Alignment(horizontal="center"))
    put(ws, f"A{r+2}", "Decision", F_BOLD, border=True)
    put(ws, f"B{r+2}", "PENDING", F_BOLD, border=True, align=Alignment(horizontal="center"))
    put(ws, f"A{r+3}", "Values are what gate_b() measured in the run (blue). Thresholds are the fund's policy, editable in "
        "yellow: change one and the Result recalculates. 'Advisory' criteria are reported but do not block. "
        "Decision stays PENDING until a named human records one on the Lightyear page.", F_NOTE)
    legend(ws, r + 5)
    widths(ws, [52, 26, 18, 14, 20, 22])
    _line_chart(f"Growth of 100: {s_name} vs {b_name}", wr, 4, 5, last, "H6", ws, "0", "Value (100 = start)")
    _line_chart("Drawdown from peak", wr, 6, 7, last, "H24", ws, "0%", "Drawdown")

    used = {"summary", "returns", "validation", "notes"}

    # ---------------------------------------------------------------- Comparators (daily)
    if comparators is not None and len(comparators.columns) > 1:
        cmp = comparators.set_index("date").reindex(df["date"].values)
        names_c = list(cmp.columns)
        wc = wb.create_sheet("Comparators")
        used.add("comparators")
        put(wc, "A1", "Simpler alternatives, run over the same days", F_TITLE)
        put(wc, "A2", "What the strategy has to beat to be worth its complexity. Daily returns (blue) from comparators.csv; "
            "values, drawdowns and the metrics table are formulas. Appendix only: the headline stays strategy vs NIFTY 500.", F_NOTE)
        header(wc, 4, ["Comparator", "CAGR", "Annualised volatility", "Sharpe vs hurdle", "Maximum drawdown",
                       "CAGR minus strategy"])
        top = 5 + len(names_c) + 2           # data header row
        d0, d1 = top + 2, top + 1 + n        # start row top+1, data rows top+2 .. top+1+n
        hdr = ["Date"]
        for c in names_c:
            hdr += [f"{c} daily return", f"{c} value", f"{c} drawdown"]
        header(wc, top, hdr, height=44)
        wc.cell(row=top + 1, column=1, value=start_date(df["date"]).to_pydatetime()).number_format = DATE
        for j in range(len(names_c)):
            cv = wc.cell(row=top + 1, column=3 + 3 * j, value=100)
            cv.font = F_INPUT
            wc.cell(row=top + 1, column=4 + 3 * j, value=0).number_format = PCT
        for i, dte in enumerate(df["date"], start=d0):
            wc.cell(row=i, column=1, value=dte.to_pydatetime()).number_format = DATE
            for j, c in enumerate(names_c):
                rc, vc, dc = 2 + 3 * j, 3 + 3 * j, 4 + 3 * j
                L = lambda k: wc.cell(row=1, column=k).column_letter
                v = cmp[c].iloc[i - d0]
                cell = wc.cell(row=i, column=rc, value=0.0 if pd.isna(v) else float(v))
                cell.font, cell.number_format = F_INPUT, RET_FMT
                wc.cell(row=i, column=vc, value=f"={L(vc)}{i-1}*(1+{L(rc)}{i})").number_format = "0.00"
                wc.cell(row=i, column=dc, value=f"={L(vc)}{i}/MAX({L(vc)}${top+1}:{L(vc)}{i})-1").number_format = PCT
        for j, c in enumerate(names_c):
            row = 5 + j
            L = lambda k: wc.cell(row=1, column=k).column_letter
            rc, vc, dc = L(2 + 3 * j), L(3 + 3 * j), L(4 + 3 * j)
            put(wc, f"A{row}", c, F_BASE, border=True)
            put(wc, f"B{row}", f"=({vc}{d1}/100)^(1/Summary!$B$9)-1", F_LINK, PCT, border=True)
            put(wc, f"C{row}", f"=STDEV({rc}{d0}:{rc}{d1})*SQRT(252)", F_BASE, PCT, border=True)
            put(wc, f"D{row}", f"=(B{row}-Hurdle)/C{row}", F_LINK, NUM2, border=True)
            put(wc, f"E{row}", f"=-MIN({dc}{d0-1}:{dc}{d1})", F_BASE, PCT, border=True)
            put(wc, f"F{row}", f"=B{row}-Summary!$B$12", F_LINK, PCT, border=True)
        sign_colour(wc, f"F5:F{4 + len(names_c)}")
        wc.freeze_panes = wc.cell(row=top + 1, column=2)
        widths(wc, [34] + [18] * (3 * len(names_c)))

    # ---------------------------------------------------------------- Holdings, Rebalances, Trades
    if holdings is not None and len(holdings):
        h = holdings.copy()
        h["date"] = pd.to_datetime(h["date"])
        if "leg" not in h.columns:
            h["leg"] = h["weight"].apply(lambda w: "long" if w > 0 else "short")
        h = h.sort_values(["date", "leg", "weight"], ascending=[True, True, False]).reset_index(drop=True)
        wh = wb.create_sheet("Holdings")
        used.add("holdings")
        extra = [c for c in h.columns if c not in ("date", "symbol", "leg", "weight")]
        put(wh, "A1", "Every position at every rebalance", F_TITLE)
        put(wh, "A2", "Target weights as a fraction of capital (shorts negative), from holdings.csv written by the run. "
            "Between rebalances the book drifts with prices.", F_NOTE)
        header(wh, 4, ["Rebalance date", "Symbol", "Leg", "Target weight"] + extra)
        h0 = 5
        for i, row in enumerate(h.itertuples(index=False), start=h0):
            put(wh, f"A{i}", row.date.to_pydatetime(), F_INPUT, DATE)
            put(wh, f"B{i}", str(row.symbol), F_INPUT)
            put(wh, f"C{i}", str(row.leg), F_INPUT)
            put(wh, f"D{i}", float(row.weight), F_INPUT, PCT)
            for k, col in enumerate(extra):
                v = _num_or_text(getattr(row, col))
                c = wh.cell(row=i, column=5 + k, value=v)
                c.font = F_INPUT
                if isinstance(v, float):
                    c.number_format = NUM2
        h1 = h0 + len(h) - 1
        sign_colour(wh, f"D{h0}:D{h1}")
        wh.freeze_panes = "A5"
        wh.auto_filter.ref = f"A4:{wh.cell(row=4, column=4 + len(extra)).column_letter}{h1}"
        widths(wh, [14, 16, 9, 14] + [16] * len(extra))
        name(wb, "H_Date", f"Holdings!$A${h0}:$A${h1}")
        name(wb, "H_Leg", f"Holdings!$C${h0}:$C${h1}")
        name(wb, "H_Weight", f"Holdings!$D${h0}:$D${h1}")

        wb2 = wb.create_sheet("Rebalances")
        used.add("rebalances")
        put(wb2, "A1", "What the strategy held at each rebalance", F_TITLE)
        put(wb2, "A2", "Counts and gross exposures are formulas over the Holdings sheet.", F_NOTE)
        header(wb2, 4, ["Rebalance date", "Names long", "Names short", "Gross long", "Gross short", "Net exposure",
                        "Gross exposure", "Largest single weight"])
        for i, dte in enumerate(sorted(h["date"].unique()), start=5):
            put(wb2, f"A{i}", pd.Timestamp(dte).to_pydatetime(), F_INPUT, DATE, border=True)
            put(wb2, f"B{i}", f'=COUNTIFS(H_Date,A{i},H_Leg,"long")', F_LINK, INT, border=True)
            put(wb2, f"C{i}", f'=COUNTIFS(H_Date,A{i},H_Leg,"short")', F_LINK, INT, border=True)
            put(wb2, f"D{i}", f'=SUMIFS(H_Weight,H_Date,A{i},H_Leg,"long")', F_LINK, PCT, border=True)
            put(wb2, f"E{i}", f'=-SUMIFS(H_Weight,H_Date,A{i},H_Leg,"short")', F_LINK, PCT, border=True)
            put(wb2, f"F{i}", f"=D{i}-E{i}", F_BASE, PCT, border=True)
            put(wb2, f"G{i}", f"=D{i}+E{i}", F_BASE, PCT, border=True)
            put(wb2, f"H{i}", f"=SUMPRODUCT(MAX((H_Date=A{i})*ABS(H_Weight)))", F_LINK, PCT, border=True)
        r_last = 4 + h["date"].nunique()
        wb2.conditional_formatting.add(f"B5:B{r_last}", CellIsRule(operator="lessThan", formula=["5"], fill=FILL_ADV))
        wb2.conditional_formatting.add(f"H5:H{r_last}", CellIsRule(operator="greaterThan", formula=["0.25"], fill=FILL_FAIL))
        put(wb2, f"A{r_last + 2}", "Amber: fewer than 5 names long. Red: one position above 25% of capital (concentration risk).", F_NOTE)
        widths(wb2, [16, 12, 12, 12, 12, 13, 14, 20])
        if trades is None:
            trades = derive_trades(h)
            t_note = "Derived by Lightyear from the change in target weights between rebalances (the run wrote no trades.csv)."
        else:
            t_note = "From trades.csv written by the run."
    else:
        t_note = "From trades.csv written by the run."

    if trades is not None and len(trades):
        t = trades.copy()
        t["date"] = pd.to_datetime(t["date"])
        wt = wb.create_sheet("Trades")
        used.add("trades")
        put(wt, "A1", "Every trade", F_TITLE)
        has_px = "price_date" in t.columns
        put(wt, "A2", t_note + " Weights are fractions of capital; shorts negative. Change = after minus before (formula)."
            + (" To check a trade: open the NSE bhavcopy for the 'Price date', find the symbol, and compare its CLOSE with "
               "'Price used'. Trades are done at the same close the ranking used (signal date = execution date), except a "
               "sale that could not trade that day, which goes at its next traded close." if has_px else ""), F_NOTE)
        wt.row_dimensions[2].height = 40
        wt["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        cols = ["date", "symbol", "action", "weight_before", "weight_after"]
        extra = [c for c in t.columns if c not in cols]
        chk = "shares" in extra and "price" in extra
        header(wt, 4, ["Trade date", "Symbol", "Action", "Weight before", "Weight after", "Change"]
               + [TRADE_HEADERS.get(c, c) for c in extra] + (["Check: shares x price (Rs)"] if chk else []), height=40)
        num_fmt = {"price": NUM2, "shares": INT, "value_rs": INT}
        for i, row in enumerate(t.itertuples(index=False), start=5):
            put(wt, f"A{i}", row.date.to_pydatetime(), F_INPUT, DATE)
            put(wt, f"B{i}", str(row.symbol), F_INPUT)
            put(wt, f"C{i}", str(row.action), F_INPUT)
            for col_l, v in (("D", row.weight_before), ("E", row.weight_after)):
                v = float(v) if pd.notna(v) else None
                put(wt, f"{col_l}{i}", v, F_INPUT, PCT)
            put(wt, f"F{i}", f'=IF(OR(D{i}="",E{i}=""),"",E{i}-D{i})', F_BASE, PCT)
            for k, col in enumerate(extra):
                raw = getattr(row, col)
                if col.endswith("date") and pd.notna(raw) and str(raw):
                    v, fmt = pd.Timestamp(raw).to_pydatetime(), DATE
                else:
                    v = _num_or_text(raw)
                    v = "" if isinstance(v, float) and v != v else v
                    fmt = num_fmt.get(col, NUM2) if isinstance(v, float) else None
                put(wt, f"{wt.cell(row=i, column=7 + k).column_letter}{i}", v, F_INPUT, fmt)
            if chk:
                L = lambda nm: wt.cell(row=1, column=7 + extra.index(nm)).column_letter
                put(wt, f"{wt.cell(row=1, column=7 + len(extra)).column_letter}{i}",
                    f'=IF(OR({L("shares")}{i}="",{L("price")}{i}=""),"",{L("shares")}{i}*{L("price")}{i})', F_BASE, INT)
        t1 = 4 + len(t)
        sign_colour(wt, f"F5:F{t1}")
        wt.freeze_panes = "D5"
        wt.auto_filter.ref = f"A4:{wt.cell(row=4, column=6 + len(extra) + (1 if chk else 0)).column_letter}{t1}"
        put(wt, "H3", "Total trades", F_BOLD)
        put(wt, "I3", f"=COUNTA(B5:B{t1})", F_BASE, INT)
        name(wb, "T_Date", f"Trades!$A$5:$A${t1}")
        name(wb, "T_Action", f"Trades!$C$5:$C${t1}")
        wide = {"why": 90, "price_basis": 44, "isin": 15}
        widths(wt, [12, 15, 9, 12, 12, 10] + [wide.get(c, 15) for c in extra] + ([18] if chk else []))

    # ---------------------------------------------------------------- Positives & negatives
    wp = wb.create_sheet("Positives & negatives")
    used.add("positives & negatives")
    put(wp, "A1", "What works and what doesn't", F_TITLE)
    put(wp, "A2", "Measured = calculated by Lightyear from the daily returns. From the tests = the run's own validation. "
        "Analyst read = Claude's review of all the results (an opinion, not a decision).", F_NOTE)
    header(wp, 4, ["", "Point", "Source"])
    rr = 5
    obs = observations or {}
    groups = [("+", obs.get("positives", []), "measured"), ("+", results.get("strengths", []) or [], "from the tests"),
              ("+", (review or {}).get("strengths", []), "analyst read"),
              ("-", obs.get("negatives", []), "measured"), ("-", results.get("weaknesses", []) or [], "from the tests"),
              ("-", (review or {}).get("weaknesses", []), "analyst read"),
              ("!", (review or {}).get("red_flags", []), "analyst read: check this")]
    for sign, items, src in groups:
        for text in items:
            fill = FILL_PASS if sign == "+" else (FILL_FAIL if sign == "-" else FILL_ADV)
            put(wp, f"A{rr}", {"+": "Positive", "-": "Negative", "!": "Red flag"}[sign], F_BOLD, fill=fill, border=True)
            put(wp, f"B{rr}", str(text), F_BASE, border=True, align=Alignment(wrap_text=True, vertical="top"))
            put(wp, f"C{rr}", src, F_NOTE, border=True)
            rr += 1
    if review and review.get("verdict"):
        put(wp, f"A{rr + 1}", "Analyst read", F_SUB)
        put(wp, f"B{rr + 1}", review["verdict"], F_BASE, align=Alignment(wrap_text=True, vertical="top"))
    widths(wp, [12, 110, 22])

    # ---------------------------------------------------------------- Validation + assumptions
    wv = wb.create_sheet("Validation")
    put(wv, "A1", "Validation, as computed by the run's code", F_TITLE)
    put(wv, "A2", "Blue = hardcoded from results.json; the source column names the file or function.", F_NOTE)
    header(wv, 4, ["Measure", "Value", "Source"])
    row = 5
    for v in results.get("validation", []) or []:
        val = _num_or_text(v.get("value"))
        fmt = FMT.get(v.get("format", "num"), NUM2) if isinstance(val, float) else None
        put(wv, f"A{row}", v.get("label", ""), border=True)
        put(wv, f"B{row}", val, F_INPUT, fmt, border=True)
        put(wv, f"C{row}", v.get("source", ""), F_NOTE, border=True)
        row += 1
    if results.get("assumptions"):
        row += 1
        put(wv, f"A{row}", "Assumptions used by the run", F_SUB)
        row += 1
        header(wv, row, ["Assumption", "Value", "Source / note"])
        row += 1
        for a in results["assumptions"]:
            val = _num_or_text(a.get("value"))
            put(wv, f"A{row}", a.get("label", ""), border=True)
            put(wv, f"B{row}", val, F_INPUT, FMT.get(a.get("format", "num"), NUM2) if isinstance(val, float) else None,
                FILL_KEY, border=True)
            put(wv, f"C{row}", a.get("note", ""), F_NOTE, border=True)
            row += 1
    widths(wv, [60, 16, 70])

    # ---------------------------------------------------------------- one sheet per table
    for t in results.get("tables", []) or []:
        p = t["csv"] if os.path.isabs(t["csv"]) else os.path.join(REPO, t["csv"])
        try:
            tab = pd.read_csv(p)
        except pd.errors.EmptyDataError:
            st = wb.create_sheet(_sheet_title(t.get("title", "Table"), used))
            put(st, "A1", t.get("title", ""), F_TITLE)
            put(st, "A2", f"{t['csv']} is empty: nothing to report. {t.get('note', '')}", F_NOTE)
            continue
        st = wb.create_sheet(_sheet_title(t.get("title", "Table"), used))
        put(st, "A1", t.get("title", ""), F_TITLE)
        put(st, "A2", f"Hardcoded from {t['csv']}. {t.get('note', '')}", F_NOTE)
        header(st, 4, [str(c) for c in tab.columns])
        fmts = []
        for col in tab.columns:
            s = pd.to_numeric(tab[col], errors="coerce").dropna()
            if PCT_NAME.search(str(col)) and len(s) and (s.abs() <= 5).all():
                fmts.append(PCT)
            elif len(s) and (s == s.round()).all():
                fmts.append(INT)
            else:
                fmts.append(NUM2)
        for i, rr_ in enumerate(tab.itertuples(index=False), start=5):
            for j, v in enumerate(rr_, start=1):
                val = _num_or_text(v)
                c = put(st, f"{st.cell(row=i, column=j).column_letter}{i}", val,
                        F_INPUT if isinstance(val, float) else F_BASE, border=True)
                if isinstance(val, float):
                    c.number_format = fmts[j - 1]
        widths(st, [max(14, min(40, len(str(c)) + 4)) for c in tab.columns])

    # ---------------------------------------------------------------- Monthly returns, Why these 20 (placed up front)
    wm = _monthly_sheet(wb, df, s_name, b_name, has_trades=("trades" in used))
    up_front = [wm]
    if selection is not None and len(selection):
        up_front.append(_selection_sheet(wb, selection))
    for k, sh in enumerate(up_front):
        wb.move_sheet(sh, offset=-(wb.index(sh) - (2 + k)))

    # ---------------------------------------------------------------- Notes
    wn = wb.create_sheet("Notes")
    put(wn, "A1", "Assumptions and caveats", F_TITLE)
    notes = list(results.get("notes", []) or []) + [
        "Data: NSE bhavcopy and Screener.in only, plus the NIFTY 500 index for the benchmark. Screener values are "
        "today's restated figures carried back, not point-in-time.",
        "Gate B is a human decision. Nothing in this workbook assigns one."]
    for i, t in enumerate(notes, start=3):
        put(wn, f"A{i}", t, align=Alignment(wrap_text=True, vertical="top"))
    wn.column_dimensions["A"].width = 130

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
    wb["Summary"].sheet_properties.tabColor = "1F4E79"
    wb.save(out_path)
    return out_path
