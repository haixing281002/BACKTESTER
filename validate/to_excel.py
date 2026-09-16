"""Export the backtest as a LIVE Excel workbook.

Not a dump of results -- a working model. Every NAV cell is a formula that
depends on the price cells above and beside it, so you can change a price and
watch the final NAV move, or click any cell and read exactly how it was made.

The point is auditability by someone who does not read Python. A portfolio
accountant, a risk officer or an IC member can trace one number from a raw price
to the headline Sharpe without taking anyone's word for what the code does.

The difference column is the one to look at first: Excel's NAV minus the
engine's NAV, on every trading day. It should be ~1e-15 on all of them.
"""
from __future__ import annotations

import os
from typing import List

import pandas as pd

OUT_DEFAULT = "outputs/validation/backtest_audit.xlsx"


def write_workbook(prices: pd.DataFrame, assets: List[str], rf: pd.Series,
                   engine_result, spread_bps: float, warmup: int,
                   rebal_dates, out: str = OUT_DEFAULT) -> str:
    """Build the audit workbook for the equal-weight, fully-invested run.

    Equal weight is chosen deliberately: it is the simplest strategy that still
    exercises every piece of accounting -- drift, month-end rebalancing, spread
    cost, turnover, compounding -- so the formulas stay readable while still
    reproducing a real result from the pipeline.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    n = len(assets)
    w_t = 1.0 / n
    half_spread = spread_bps / 1e4 / 2.0
    rf_d = float(rf.iloc[0])
    idx = prices.index
    nrows = len(idx)
    last = nrows + 1                       # last worksheet row holding data

    # ---- column layout, derived once so the prose can quote real letters ----
    L = {}
    order = (["date"] + [f"px{j}" for j in range(n)] + ["rebal", "started", "navpre"]
             + [f"u{j}" for j in range(n)]
             + ["cash", "traded", "cost", "nav", "ret", "peak", "dd", "eng", "diff"])
    for i, key in enumerate(order, start=1):
        L[key] = get_column_letter(i)
    PX = [L[f"px{j}"] for j in range(n)]
    UN = [L[f"u{j}"] for j in range(n)]

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2F4858")
    warn = PatternFill("solid", fgColor="7A3B2E")

    # ---------------------------------------------------------------- README
    ws = wb.active
    ws.title = "READ ME"
    lines = [
        ("What this workbook is", True),
        ("", False),
        ("A live re-computation of the equal-weight, fully-invested backtest from", False),
        ("the Research OS engine. Every NAV, cost and metric is an Excel FORMULA,", False),
        ("not a pasted number. Click any cell and read the formula bar.", False),
        ("", False),
        ("How to check the engine with it", True),
        ("", False),
        (f"1. On the 'Backtest' sheet, look at column {L['diff']} (NAV difference).", False),
        ("   That is Excel's NAV minus the engine's NAV, on every trading day.", False),
        ("   Every row should read 0.00E+00 or something around 1e-15.", False),
        ("2. The 'Metrics' sheet reports the largest difference across all rows.", False),
        (f"3. Pick any rebalance row (column {L['rebal']} = 1) and check by hand:", False),
        (f"      traded ({L['traded']}) = sum over the five sleeves of", False),
        ("                   |target weight - drifted weight|", False),
        (f"      cost   ({L['cost']}) = traded x half-spread "
         f"({spread_bps:g}bp / 2 = {half_spread:.5f})", False),
        (f"      NAV    ({L['nav']}) = NAV before cost ({L['navpre']}) x (1 - cost)", False),
        ("4. Change a price on the 'Prices' sheet. Everything downstream moves.", False),
        ("   Undo it afterwards, or the workbook no longer matches the engine.", False),
        ("", False),
        ("Why equal weight", True),
        ("", False),
        ("It is the simplest strategy that still exercises every piece of the", False),
        ("accounting: drift, month-end rebalancing, spread cost, turnover and", False),
        ("compounding. The formulas stay readable and it is still a real result", False),
        ("from the pipeline -- the benchmark the dynamic strategy failed to beat.", False),
        ("", False),
        ("What agreement proves", True),
        ("", False),
        ("That the ARITHMETIC is right: compounding, cost timing, turnover, the", False),
        ("causal lag, month-end rebalancing, and the headline metrics.", False),
        ("", False),
        ("What it does NOT prove", True),
        ("", False),
        ("That the ASSUMPTIONS are right. This workbook charges 30bp because the", False),
        ("Strategy Card says 30bp. It uses a flat 6% cash proxy because no Indian", False),
        ("risk-free series exists in the snapshot. It inherits the backfill bias in", False),
        ("the NSE factor indices -- every series starts at exactly 1000.00, which", False),
        ("means the history was reconstructed after the rules were written.", False),
        ("", False),
        ("Excel agreeing with Python about a wrong assumption is two tools being", False),
        ("precisely wrong together. Those questions belong to Gate A and Gate B,", False),
        ("and no amount of cross-checking retires them.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if bold:
            c.font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 84

    # --------------------------------------------------------------- Prices
    wsp = wb.create_sheet("Prices")
    c = wsp.cell(row=1, column=1, value="Date")
    c.font, c.fill = head, fill
    for j, a in enumerate(assets):
        c = wsp.cell(row=1, column=2 + j, value=a)
        c.font, c.fill = head, fill
    arr = prices[assets].to_numpy(dtype=float)
    for i in range(nrows):
        wsp.cell(row=2 + i, column=1,
                 value=idx[i].to_pydatetime()).number_format = "yyyy-mm-dd"
        for j in range(n):
            wsp.cell(row=2 + i, column=2 + j, value=float(arr[i, j]))
    wsp.column_dimensions["A"].width = 12
    for j in range(n):
        wsp.column_dimensions[get_column_letter(2 + j)].width = 16
    wsp.freeze_panes = "B2"

    # ------------------------------------------------------------- Backtest
    ws = wb.create_sheet("Backtest")
    titles = (["Date"] + [f"Price: {a}" for a in assets]
              + ["Rebal?", "Started", "NAV before cost"]
              + [f"Units: {a}" for a in assets]
              + ["Cash", "Traded (L1)", "Cost", "NAV", "Daily return",
                 "Peak", "Drawdown", "ENGINE NAV", "NAV difference"])
    for j, name in enumerate(titles, start=1):
        c = ws.cell(row=1, column=j, value=name)
        c.font, c.fill = head, fill
        c.alignment = Alignment(wrap_text=True, vertical="center")

    eng_nav = engine_result.value.to_numpy(dtype=float)
    col = {k: i for i, k in enumerate(order, start=1)}

    for i in range(nrows):
        r, pr = 2 + i, 1 + i
        ws.cell(row=r, column=col["date"],
                value=idx[i].to_pydatetime()).number_format = "yyyy-mm-dd"
        for j in range(n):
            ws.cell(row=r, column=col[f"px{j}"], value=f"=Prices!{PX[j]}{r}")

        # Rebal? -- the last trading day present in its calendar month, once
        # past warmup. Comparing this row's month to the NEXT row's is exactly
        # "is there no later trading day in this month".
        month_end = ("TRUE" if r == last else
                     f'TEXT({L["date"]}{r},"yyyy-mm")<>TEXT({L["date"]}{r+1},"yyyy-mm")')
        ws.cell(row=r, column=col["rebal"],
                value=f"=IF(AND({i}>={warmup},{month_end}),1,0)")

        if i == 0:
            ws.cell(row=r, column=col["started"], value=f"={L['rebal']}{r}")
            ws.cell(row=r, column=col["navpre"], value=1.0)
            for j in range(n):
                ws.cell(row=r, column=col[f"u{j}"], value=0.0)
            ws.cell(row=r, column=col["cash"], value=1.0)
            ws.cell(row=r, column=col["traded"], value=0.0)
            ws.cell(row=r, column=col["cost"], value=0.0)
            ws.cell(row=r, column=col["nav"], value=1.0)
            ws.cell(row=r, column=col["ret"], value=0.0)
            ws.cell(row=r, column=col["peak"], value=1.0)
            ws.cell(row=r, column=col["dd"], value=0.0)
        else:
            ws.cell(row=r, column=col["started"],
                    value=f"=IF(OR({L['started']}{pr}=1,{L['rebal']}{r}=1),1,0)")

            # Holdings marked to today's close plus one day of cash accrual.
            # Flat at 1.0 until the first trade, which mirrors the engine.
            held = (f"SUMPRODUCT({UN[0]}{pr}:{UN[-1]}{pr},"
                    f"{PX[0]}{r}:{PX[-1]}{r})")
            ws.cell(row=r, column=col["navpre"], value=(
                f"=IF({L['started']}{pr}=0,1,{held}+{L['cash']}{pr}*(1+{rf_d!r}))"))

            # L1 distance between target and drifted weights. On the very first
            # trade the book is empty, so the whole target is bought: L1 = 1.
            terms = "+".join(
                f"ABS({w_t!r}-{UN[j]}{pr}*{PX[j]}{r}/{L['navpre']}{r})"
                for j in range(n))
            ws.cell(row=r, column=col["traded"], value=(
                f"=IF({L['rebal']}{r}=1,"
                f"IF({L['started']}{pr}=0,1,{terms}),0)"))
            ws.cell(row=r, column=col["cost"],
                    value=f"={L['traded']}{r}*{half_spread!r}")
            ws.cell(row=r, column=col["nav"],
                    value=f"={L['navpre']}{r}*(1-{L['cost']}{r})")

            for j in range(n):
                ws.cell(row=r, column=col[f"u{j}"], value=(
                    f"=IF({L['rebal']}{r}=1,{w_t!r}*{L['nav']}{r}/{PX[j]}{r},"
                    f"{UN[j]}{pr})"))
            # Fully invested: the target weights sum to 1, so a rebalance leaves
            # no cash. Between trades whatever is there accrues.
            ws.cell(row=r, column=col["cash"], value=(
                f"=IF({L['rebal']}{r}=1,{L['nav']}{r}*(1-{float(n*w_t)!r}),"
                f"IF({L['started']}{pr}=0,1,{L['cash']}{pr}*(1+{rf_d!r})))"))

            ws.cell(row=r, column=col["ret"],
                    value=f"={L['nav']}{r}/{L['nav']}{pr}-1")
            ws.cell(row=r, column=col["peak"],
                    value=f"=MAX({L['peak']}{pr},{L['nav']}{r})")
            ws.cell(row=r, column=col["dd"],
                    value=f"={L['nav']}{r}/{L['peak']}{r}-1")

        ws.cell(row=r, column=col["eng"], value=float(eng_nav[i]))
        c = ws.cell(row=r, column=col["diff"],
                    value=f"={L['nav']}{r}-{L['eng']}{r}")
        c.number_format = "0.00E+00"

    widths = {"date": 12, "rebal": 8, "started": 9, "navpre": 17, "cash": 12,
              "traded": 13, "cost": 12, "nav": 14, "ret": 13, "peak": 14,
              "dd": 12, "eng": 14, "diff": 15}
    for key, wdt in widths.items():
        ws.column_dimensions[L[key]].width = wdt
    for c_ in PX + UN:
        ws.column_dimensions[c_].width = 15
    ws.freeze_panes = "B2"

    # -------------------------------------------------------------- Metrics
    wsm = wb.create_sheet("Metrics")

    def rng(key):
        return f"Backtest!{L[key]}2:{L[key]}{last}"

    # Parenthesised as a whole term. Without the outer brackets, "1/{years}"
    # expands to 1/(days)/365.25 -- which Excel reads left to right as
    # (1/days)/365.25, not 1/(days/365.25). That silently turned the CAGR into
    # a number near zero, and it is exactly what the formula evaluator caught.
    years = f"((Backtest!{L['date']}{last}-Backtest!{L['date']}2)/365.25)"
    nav_last = f"Backtest!{L['nav']}{last}"
    nav_first = f"Backtest!{L['nav']}2"
    cash_cagr = f"((1+{rf_d!r})^({nrows}/{years})-1)"

    rows = [
        ("Metric", "Excel (live formula)", "Note"),
        ("Final NAV", f"={nav_last}", "grew from 1.00"),
        ("CAGR", f"=({nav_last}/{nav_first})^(1/{years})-1", ""),
        ("Annualised volatility", f"=STDEV.S({rng('ret')})*SQRT(252)", "sample sd x sqrt(252)"),
        ("Max drawdown", f"=-MIN({rng('dd')})", ""),
        ("Compounded cash return (CAGR)", f"={cash_cagr}", "flat proxy -- see READ ME"),
        ("Sharpe (geometric, paper definition)", "=(B3-B6)/B4", "(CAGR - cash) / vol"),
        ("Sharpe (conventional)",
         f"=(AVERAGE({rng('ret')})-{rf_d!r})/STDEV.S({rng('ret')})*SQRT(252)",
         "mean excess / sd"),
        ("Total cost paid (fraction of NAV)", f"=SUM({rng('cost')})", ""),
        ("Total turnover (one-way)", f"=SUM({rng('traded')})/2", ""),
        ("Number of rebalances", f"=SUM({rng('rebal')})", ""),
        ("", "", ""),
        ("LARGEST NAV DIFFERENCE vs engine",
         # MAX/-MIN rather than MAX(ABS(...)), which needs array entry in
         # Excel versions older than 365.
         f"=MAX(MAX({rng('diff')}),-MIN({rng('diff')}))",
         "should be ~1e-15"),
    ]
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row, start=1):
            c = wsm.cell(row=i, column=j, value=v)
            if i == 1:
                c.font, c.fill = head, fill
    wsm.cell(row=13, column=1).fill = warn
    wsm.cell(row=13, column=1).font = Font(bold=True, color="FFFFFF")
    wsm.cell(row=13, column=2).number_format = "0.00E+00"
    wsm.column_dimensions["A"].width = 38
    wsm.column_dimensions["B"].width = 46
    wsm.column_dimensions["C"].width = 24

    wb.save(out)
    return out
