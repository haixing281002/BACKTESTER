"""A second, independent implementation of the backtest accounting.

WHY THIS FILE EXISTS

`ros/engine/backtest.py` is tested, but every one of those tests was written by
the same author as the engine. A shared misunderstanding -- about what a weight
means after a day of drift, about when a cost is charged, about which bar a
signal may be traded on -- would be baked into the engine AND into its tests,
and the suite would pass while the numbers were wrong.

This file exists to make that failure mode detectable. It computes the same
quantities a second time, and it is built to disagree if the first one is wrong:

  1. It imports NOTHING from `ros`. That is enforced by a test
     (tests/test_independent_agrees.py), not by good intentions.

  2. It uses a DIFFERENT FORMULATION. The engine carries portfolio WEIGHTS and
     renormalises them each day. This file carries UNIT HOLDINGS and a CASH
     BALANCE, and derives NAV as units . prices + cash -- the way a fund
     accountant would. The two are mathematically equivalent, so they must
     agree; but a renormalisation bug lives in one and not the other.

  3. It reads the workbook through raw openpyxl rather than pandas.read_excel,
     so a parsing difference shows up as a data mismatch rather than being
     inherited silently.

Agreement here is not proof of correctness -- both could be wrong about
something neither implementation questions, and no amount of cross-checking
fixes a bad assumption. What agreement DOES rule out is an arithmetic or
bookkeeping slip in the engine, which is the most common way a backtest lies.
Where the two must agree by construction, the tolerance is set at 1e-10.
"""
from __future__ import annotations

import math
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

ANN = 252


# ---------------------------------------------------------------------------
# 1. Data, read a different way
# ---------------------------------------------------------------------------
def read_workbook(path: str) -> pd.DataFrame:
    """Parse the NSE factor workbook with openpyxl directly.

    Deliberately not pandas.read_excel: if the two parsers disagree about where
    the header sits or which column is which, that is a finding, and it can only
    surface if the second reader is genuinely a second reader.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    date_row = date_col = None
    for r in range(min(12, len(rows))):
        for c, v in enumerate(rows[r]):
            if isinstance(v, str) and v.strip().lower() == "date":
                date_row, date_col = r, c
                break
        if date_row is not None:
            break
    if date_row is None:
        raise ValueError(f"{path}: no 'Date' header cell found")

    name_row = date_row - 1
    body = rows[date_row + 1:]
    dates = pd.to_datetime([r[date_col] if date_col < len(r) else None for r in body],
                           errors="coerce")

    cols: Dict[str, List[Optional[float]]] = {}
    width = max(len(r) for r in rows[:date_row + 1])
    for c in range(width):
        if c == date_col:
            continue
        nm = rows[name_row][c] if name_row >= 0 and c < len(rows[name_row]) else None
        if not isinstance(nm, str) or not nm.strip():
            continue
        vals = []
        for r in body:
            v = r[c] if c < len(r) else None
            vals.append(float(v) if isinstance(v, (int, float)) and not isinstance(v, bool)
                        else None)
        if any(v is not None for v in vals):
            cols[nm.strip()] = vals

    frame = pd.DataFrame(cols, index=dates)
    frame = frame[frame.index.notna()].sort_index()
    return frame[~frame.index.duplicated(keep="last")]


# ---------------------------------------------------------------------------
# 2. Metrics, computed by hand
# ---------------------------------------------------------------------------
def ann_vol(returns: Sequence[float], ann: int = ANN) -> float:
    r = np.asarray([x for x in returns if np.isfinite(x)], dtype=float)
    if len(r) < 3:
        return float("nan")
    # explicit two-pass sample variance rather than pandas .std(ddof=1)
    mu = r.sum() / len(r)
    var = ((r - mu) ** 2).sum() / (len(r) - 1)
    return math.sqrt(var) * math.sqrt(ann)


def cagr(value: pd.Series) -> float:
    v = value.dropna()
    if len(v) < 2 or v.iloc[0] <= 0:
        return float("nan")
    years = (v.index[-1] - v.index[0]).days / 365.25
    return float(v.iloc[-1] / v.iloc[0]) ** (1.0 / years) - 1.0 if years > 0 else float("nan")


def max_drawdown(value: pd.Series) -> float:
    """Running-peak loop, not value/value.cummax().

    Same answer, different code path -- which is the entire point.
    """
    peak = -np.inf
    worst = 0.0
    for v in value.to_numpy(dtype=float):
        if v > peak:
            peak = v
        if peak > 0:
            dd = v / peak - 1.0
            if dd < worst:
                worst = dd
    return -worst


def compounded_cash_cagr(rf_daily: Sequence[float], index: pd.DatetimeIndex) -> float:
    years = (index[-1] - index[0]).days / 365.25
    if years <= 0:
        return float("nan")
    total = 1.0
    for x in rf_daily:                      # sequential product, not np.prod
        total *= (1.0 + x)
    return total ** (1.0 / years) - 1.0


def sharpe_geometric(value: pd.Series, returns: Sequence[float],
                     rf_daily: Optional[Sequence[float]] = None) -> float:
    v = ann_vol(returns)
    if not np.isfinite(v) or v == 0:
        return float("nan")
    rf_c = compounded_cash_cagr(rf_daily, value.index) if rf_daily is not None else 0.0
    return (cagr(value) - rf_c) / v


def sharpe_conventional(returns: Sequence[float],
                        rf_daily: Optional[Sequence[float]] = None) -> float:
    r = np.asarray(returns, dtype=float)
    keep = np.isfinite(r)
    r = r[keep]
    if len(r) < 3:
        return float("nan")
    if rf_daily is not None:
        r = r - np.asarray(rf_daily, dtype=float)[keep]
    mu = r.sum() / len(r)
    var = ((r - mu) ** 2).sum() / (len(r) - 1)
    sd = math.sqrt(var)
    return float(mu / sd * math.sqrt(ANN)) if sd > 0 else float("nan")


# ---------------------------------------------------------------------------
# 3. Signals
# ---------------------------------------------------------------------------
def cash_series(index: pd.DatetimeIndex, annual_rate: float) -> pd.Series:
    """Constant declared proxy. (1+annual)^(1/252) - 1, per trading day."""
    return pd.Series((1.0 + annual_rate) ** (1.0 / 252.0) - 1.0, index=index)


def trailing_portfolio_vol(returns: pd.DataFrame, weights: np.ndarray,
                           window: int) -> pd.Series:
    """Annualised trailing vol of a fixed-weight mix, full windows only.

    Written as an explicit window loop rather than .rolling(...).std(), so an
    off-by-one in the window boundary cannot be shared with the engine.
    """
    port = returns.to_numpy(dtype=float) @ np.asarray(weights, dtype=float)
    out = np.full(len(port), np.nan)
    for i in range(window - 1, len(port)):
        block = port[i - window + 1:i + 1]
        if np.isnan(block).any():
            continue
        out[i] = ann_vol(block)
    return pd.Series(out, index=returns.index)


def month_end_dates(index: pd.DatetimeIndex) -> List[pd.Timestamp]:
    """Last TRADING day present in each calendar month."""
    seen: Dict[Tuple[int, int], pd.Timestamp] = {}
    for d in index:
        seen[(d.year, d.month)] = d          # later date overwrites earlier
    return sorted(seen.values())


# ---------------------------------------------------------------------------
# 4. The backtest, as a fund accountant would keep it
# ---------------------------------------------------------------------------
class UnitsBacktest:
    """Unit holdings and a cash balance. NAV = units . prices + cash.

    The engine tracks weights and renormalises them every day. This tracks what
    a custodian would track. Both must produce the same NAV path; only one of
    them can have a renormalisation bug.
    """

    def __init__(self, prices: pd.DataFrame, assets: List[str],
                 rf_daily: pd.Series, spread_bps: float, lag_days: int,
                 allow_cash: bool):
        self.assets = list(assets)
        self.prices = prices[self.assets].astype(float)
        if self.prices.isna().any().any():
            raise ValueError("price frame has NaNs for the investable set")
        self.rf = rf_daily.reindex(self.prices.index).fillna(0.0)
        self.half_spread = float(spread_bps) / 1e4 / 2.0
        self.lag_days = int(lag_days)
        self.allow_cash = bool(allow_cash)
        self.n = len(self.assets)

    def run(self, target_fn, rebalance_on: Sequence[pd.Timestamp],
            warmup: int = 0, name: str = "independent") -> Dict[str, object]:
        idx = self.prices.index
        P = self.prices.to_numpy(dtype=float)
        rf = self.rf.to_numpy(dtype=float)
        rebal = set(pd.DatetimeIndex(rebalance_on))

        units = np.zeros(self.n)
        cash = 1.0
        started = False                      # mirrors the engine: flat until traded

        navs, turns, costs, cashw = [], [], [], []
        w = np.zeros(self.n)
        n_rebals = 0

        for i, date in enumerate(idx):
            # --- drift ----------------------------------------------------
            if started:
                cash *= (1.0 + rf[i])
                nav = float(units @ P[i]) + cash
            else:
                nav = 1.0                    # nothing held, nothing accrues yet

            day_turn = 0.0
            day_cost = 0.0

            # --- rebalance ------------------------------------------------
            if date in rebal and i >= warmup:
                held_value = units * P[i]
                w_now = held_value / nav if nav > 0 else np.zeros(self.n)
                target = target_fn(date, i, w_now)
                if target is not None:
                    target = np.clip(np.asarray(target, dtype=float), 0.0, None)
                    if not self.allow_cash:
                        tot = target.sum()
                        target = (target / tot) if tot > 0 else np.ones(self.n) / self.n
                    if target.sum() > 1.0 + 1e-9:
                        target = target / target.sum()

                    traded = float(np.abs(target - w_now).sum())
                    day_cost = self.half_spread * traded
                    day_turn = 0.5 * traded
                    nav *= (1.0 - day_cost)          # cost paid out of NAV
                    units = target * nav / P[i]      # re-establish holdings
                    cash = nav * (1.0 - target.sum())
                    w = target
                    started = True
                    n_rebals += 1

            if started:
                w = (units * P[i]) / nav if nav > 0 else np.zeros(self.n)

            navs.append(nav)
            turns.append(day_turn)
            costs.append(day_cost)
            cashw.append(1.0 - float(w.sum()))

        value = pd.Series(navs, index=idx, name=name)
        return {
            "value": value,
            "returns": value.pct_change().fillna(0.0),
            "turnover": pd.Series(turns, index=idx),
            "costs": pd.Series(costs, index=idx),
            "cash_weight": pd.Series(cashw, index=idx),
            "n_rebalances": n_rebals,
        }


# ---------------------------------------------------------------------------
# 5. Allocators, re-derived from the card and the paper
# ---------------------------------------------------------------------------
def fixed_weight_fn(w: np.ndarray):
    return lambda date, i, w_now: w.copy()


def vol_target_fn(w: np.ndarray, target_vol: float, sigma: pd.Series,
                  max_scale: float = 1.0):
    """Paper eq. (1): dilute a fixed mix with cash when trailing vol is too high.

        scale = min(max_scale, target_vol / sigma_t)

    `sigma` must ALREADY be lag-shifted by the caller -- see shift_causal.
    """
    def fn(date, i, w_now):
        s = sigma.iloc[i]
        if not np.isfinite(s) or s <= 0:
            return w.copy()
        return w * min(max_scale, target_vol / float(s))
    return fn


def shift_causal(obj, lag_days: int):
    """Shift by 1 + lag_days.

    The +1 forbids trading on the bar whose close produced the signal; lag_days
    is the card's declared implementation lag on top of that. NSE index closes
    publish after the close, which is why the card sets lag_days >= 1.
    """
    return obj.shift(1 + int(lag_days))
